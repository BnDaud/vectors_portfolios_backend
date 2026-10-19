import importlib
import zlib
import io
import json
from datetime import date, timedelta
from unittest import mock

from django.apps import apps as django_apps
from django.contrib.auth.models import User
from django.core.cache import cache
from django.core.management import call_command
from io import StringIO
from django.test import SimpleTestCase, TransactionTestCase, override_settings
import uuid
from datetime import timedelta
from django.utils import timezone
from rest_framework.throttling import ScopedRateThrottle
from rest_framework.test import APIClient, APITestCase

from .models import (
    About, ENTRY_DESCRIPTION_MAX, PageView, Portfolio, PortfolioTrack, Profile, Proficiency, Resume, ResumeView,
    School, TRACK_CONTACT_FIELDS, Visitor,
)

API = "/api_root"


def make_project(track, **fields):
    """A project shown in `track` (more portfolios can be added with .tracks.add)."""
    project = Portfolio.objects.create(**fields)
    project.tracks.add(track)
    return project


class TrackIndependenceTests(APITestCase):
    """Every portfolio (track) owns its dp, links and about; duplicating a
    track copies them once and the two evolve separately afterwards."""

    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user("owner", password="pw", first_name="O", last_name="W")
        cls.profile = Profile.objects.create(
            user=cls.user,
            display_pic="https://img.example.com/profile.png",
            github="https://github.com/owner",
            linkedin="https://linkedin.com/in/owner",
            whatapp="https://wa.me/111",
            facebook="https://facebook.com/owner",
            twitter="https://x.com/owner",
            email="owner@example.com",
            phone_number="+111",
        )
        cls.track_a = cls.make_track("Full Stack", "full-stack", dp="https://img.example.com/a.png")

    @classmethod
    def make_track(cls, name, slug, dp):
        track = PortfolioTrack.objects.create(
            profile=cls.profile, name=name, slug=slug, display_pic=dp,
            github="https://github.com/a", email="a@example.com",
        )
        about = About.objects.create(
            track=track, skill=f"{name} dev", description=f"about {name}",
            image_link="https://img.example.com/about-a.png", experience_since=date(2022, 1, 1),
        )
        Resume.objects.create(about=about, post="Dev", start_year=date(2022, 1, 1),
                              end_year=date(2023, 1, 1), company_name="Acme")
        School.objects.create(about=about, school_name="Uni", start_year=date(2018, 1, 1),
                              end_year=date(2022, 1, 1))
        Proficiency.objects.create(about=about, skill_name="Python", skill_range=80)
        make_project(track, category="WebDev", name="Site",
                                 thumbnail="https://a.com/t.png", project_link="https://a.com")
        return track

    def setUp(self):
        self.client = APIClient(HTTP_HOST="localhost")
        self.client.force_authenticate(user=self.user)

    def duplicate(self, source, name="Copy", copy_items=True):
        res = self.client.post(f"{API}/tracks/", {
            "username": "owner", "name": name, "copy_from": source.id,
            "copy_portfolio_items": copy_items,
        }, format="json")
        self.assertEqual(res.status_code, 201, res.content)
        return PortfolioTrack.objects.get(pk=res.json()["id"])

    # --- independence between tracks -------------------------------------

    def test_editing_one_tracks_dp_leaves_others_and_profile_untouched(self):
        track_b = self.make_track("CAD", "cad", dp="https://img.example.com/b.png")

        res = self.client.patch(f"{API}/tracks/{self.track_a.id}",
                                {"display_pic": "https://img.example.com/new-a.png"}, format="json")
        self.assertEqual(res.status_code, 200, res.content)

        self.track_a.refresh_from_db(); track_b.refresh_from_db(); self.profile.refresh_from_db()
        self.assertEqual(self.track_a.display_pic, "https://img.example.com/new-a.png")
        self.assertEqual(track_b.display_pic, "https://img.example.com/b.png")
        self.assertEqual(self.profile.display_pic, "https://img.example.com/profile.png")

    def test_every_contact_field_is_editable_per_track(self):
        track_b = self.make_track("CAD", "cad", dp="https://img.example.com/b.png")
        payload = {
            "display_pic": "https://img.example.com/x.png", "github": "https://github.com/x",
            "linkedin": "https://linkedin.com/in/x", "whatapp": "https://wa.me/222",
            "facebook": "https://facebook.com/x", "twitter": "https://x.com/x",
            "email": "x@example.com", "phone_number": "+222",
        }
        self.assertEqual(set(payload), set(TRACK_CONTACT_FIELDS))
        res = self.client.patch(f"{API}/tracks/{self.track_a.id}", payload, format="json")
        self.assertEqual(res.status_code, 200, res.content)

        self.track_a.refresh_from_db(); track_b.refresh_from_db()
        for field, value in payload.items():
            self.assertEqual(getattr(self.track_a, field), value, field)
            self.assertNotEqual(getattr(track_b, field), value, field)

    def test_about_is_per_track(self):
        track_b = self.make_track("CAD", "cad", dp="https://img.example.com/b.png")
        res = self.client.patch(f"{API}/about/{self.track_a.about.id}",
                                {"description": "changed", "skill": "Changed skill"}, format="json")
        self.assertEqual(res.status_code, 200, res.content)
        track_b.about.refresh_from_db()
        self.assertEqual(track_b.about.description, "about CAD")
        self.assertEqual(track_b.about.skill, "CAD dev")

    def test_tracks_endpoint_exposes_dp_for_each_landing_tile(self):
        track_b = self.make_track("CAD", "cad", dp="https://img.example.com/b.png")
        data = self.client.get(f"{API}/profiles/?username=owner").json()[0]
        dps = {t["id"]: t["display_pic"] for t in data["tracks"]}
        self.assertEqual(dps, {self.track_a.id: "https://img.example.com/a.png",
                               track_b.id: "https://img.example.com/b.png"})

    # --- duplicating -----------------------------------------------------

    def test_duplicate_carries_all_content_over(self):
        copy = self.duplicate(self.track_a)

        self.assertNotEqual(copy.id, self.track_a.id)
        self.assertNotEqual(copy.slug, self.track_a.slug)
        for field in TRACK_CONTACT_FIELDS:
            self.assertEqual(getattr(copy, field), getattr(self.track_a, field), field)
        self.assertEqual(copy.about.skill, self.track_a.about.skill)
        self.assertEqual(copy.about.description, self.track_a.about.description)
        self.assertEqual(copy.about.image_link, self.track_a.about.image_link)
        self.assertEqual(copy.about.experience_since, self.track_a.about.experience_since)
        self.assertEqual(copy.about.resume.count(), 1)
        self.assertEqual(copy.about.school.count(), 1)
        self.assertEqual(copy.about.proficiency.count(), 1)
        self.assertEqual(copy.items.count(), 1)
        self.assertEqual(copy.items.get().name, "Site")

    def test_duplicate_can_skip_portfolio_items(self):
        copy = self.duplicate(self.track_a, copy_items=False)
        self.assertEqual(copy.items.count(), 0)
        self.assertEqual(copy.about.resume.count(), 1)

    def test_editing_the_copy_never_changes_the_original(self):
        copy = self.duplicate(self.track_a)
        self.client.patch(f"{API}/tracks/{copy.id}",
                          {"display_pic": "https://img.example.com/copy.png", "github": ""}, format="json")
        self.client.patch(f"{API}/about/{copy.about.id}", {"description": "copy text"}, format="json")
        self.client.patch(f"{API}/resume/{copy.about.resume.get().id}", {"post": "Copy post"}, format="json")
        self.client.patch(f"{API}/portfolio/{copy.items.get().id}", {"name": "Copy site"}, format="json")

        self.track_a.refresh_from_db()
        self.assertEqual(self.track_a.display_pic, "https://img.example.com/a.png")
        self.assertEqual(self.track_a.github, "https://github.com/a")
        self.assertEqual(self.track_a.about.description, "about Full Stack")
        self.assertEqual(self.track_a.about.resume.get().post, "Dev")
        self.assertEqual(self.track_a.items.get().name, "Site")

    def test_editing_the_original_never_changes_the_copy(self):
        copy = self.duplicate(self.track_a)
        self.client.patch(f"{API}/tracks/{self.track_a.id}",
                          {"display_pic": "https://img.example.com/orig-new.png"}, format="json")
        self.client.patch(f"{API}/about/{self.track_a.about.id}", {"description": "orig text"}, format="json")

        copy.refresh_from_db()
        self.assertEqual(copy.display_pic, "https://img.example.com/a.png")
        self.assertEqual(copy.about.description, "about Full Stack")

    def test_deleting_the_copy_leaves_the_original_intact(self):
        copy = self.duplicate(self.track_a)
        res = self.client.delete(f"{API}/tracks/{copy.id}")
        self.assertEqual(res.status_code, 204)
        self.track_a.refresh_from_db()
        self.assertEqual(self.track_a.about.resume.count(), 1)
        self.assertEqual(self.track_a.items.count(), 1)

    def test_blank_track_starts_from_the_owner_profile_then_stands_alone(self):
        res = self.client.post(f"{API}/tracks/", {"username": "owner", "name": "Blank"}, format="json")
        self.assertEqual(res.status_code, 201, res.content)
        blank = PortfolioTrack.objects.get(pk=res.json()["id"])
        for field in TRACK_CONTACT_FIELDS:
            self.assertEqual(getattr(blank, field), getattr(self.profile, field), field)
        self.assertEqual(blank.about.skill, "")

        self.client.patch(f"{API}/tracks/{blank.id}", {"display_pic": "https://img.example.com/blank.png"},
                          format="json")
        self.profile.refresh_from_db()
        self.assertEqual(self.profile.display_pic, "https://img.example.com/profile.png")

    def test_bad_copy_from_is_rejected_and_creates_nothing(self):
        before = PortfolioTrack.objects.count()
        res = self.client.post(f"{API}/tracks/", {"username": "owner", "name": "X", "copy_from": 99999},
                               format="json")
        self.assertEqual(res.status_code, 400)
        self.assertEqual(PortfolioTrack.objects.count(), before)

    def test_cannot_copy_from_another_owners_track(self):
        other_user = User.objects.create_user("other", password="pw")
        other_profile = Profile.objects.create(user=other_user, display_pic="https://img.example.com/o.png")
        foreign = PortfolioTrack.objects.create(profile=other_profile, name="Theirs", slug="theirs")
        before = PortfolioTrack.objects.count()
        res = self.client.post(f"{API}/tracks/", {"username": "owner", "name": "X", "copy_from": foreign.id},
                               format="json")
        self.assertEqual(res.status_code, 400)
        self.assertEqual(PortfolioTrack.objects.count(), before)

    def test_duplicating_twice_gives_unique_slugs(self):
        first = self.duplicate(self.track_a, name="Full Stack")
        second = self.duplicate(self.track_a, name="Full Stack")
        self.assertEqual(len({self.track_a.slug, first.slug, second.slug}), 3)

    # --- validation / auth -----------------------------------------------

    def test_track_dp_cannot_be_blanked_or_set_to_a_non_url(self):
        for payload in ({"display_pic": ""}, {"display_pic": "not a url"}, {"github": "github.com/x"},
                        {"email": "x" * 51}):
            res = self.client.patch(f"{API}/tracks/{self.track_a.id}", payload, format="json")
            self.assertEqual(res.status_code, 400, payload)
        self.track_a.refresh_from_db()
        self.assertEqual(self.track_a.display_pic, "https://img.example.com/a.png")

    def test_optional_links_can_be_cleared(self):
        res = self.client.patch(f"{API}/tracks/{self.track_a.id}",
                                {"github": "", "twitter": "", "phone_number": ""}, format="json")
        self.assertEqual(res.status_code, 200, res.content)

    def test_anonymous_cannot_edit_a_track(self):
        anon = APIClient(HTTP_HOST="localhost")
        res = anon.patch(f"{API}/tracks/{self.track_a.id}", {"display_pic": "https://x.com/y.png"},
                         format="json")
        self.assertIn(res.status_code, (401, 403))
        self.assertEqual(anon.get(f"{API}/tracks/").status_code, 200)

    def test_rename_still_works_and_keeps_contact_fields(self):
        res = self.client.patch(f"{API}/tracks/{self.track_a.id}", {"name": "Renamed"}, format="json")
        self.assertEqual(res.status_code, 200, res.content)
        self.track_a.refresh_from_db()
        self.assertEqual(self.track_a.slug, "renamed")
        self.assertEqual(self.track_a.display_pic, "https://img.example.com/a.png")


class RecentChangesTests(APITestCase):
    """CAD category and the 1000 character About description."""

    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user("owner", password="pw")
        profile = Profile.objects.create(user=cls.user, display_pic="https://img.example.com/p.png")
        cls.track = PortfolioTrack.objects.create(profile=profile, name="T", slug="t")
        cls.about = About.objects.create(track=cls.track, skill="s", description="d")

    def setUp(self):
        self.client = APIClient(HTTP_HOST="localhost")
        self.client.force_authenticate(user=self.user)

    def test_cad_category_accepted_and_wrong_case_rejected(self):
        body = {"track": self.track.id, "name": "n", "thumbnail": "https://a.com/t.png",
                "project_link": "https://a.com"}
        ok = self.client.post(f"{API}/portfolio/", {**body, "category": "CAD"}, format="json")
        self.assertEqual(ok.status_code, 201, ok.content)
        bad = self.client.post(f"{API}/portfolio/", {**body, "category": "cad"}, format="json")
        self.assertEqual(bad.status_code, 400)

    def test_description_limit_is_1000(self):
        url = f"{API}/about/{self.about.id}"
        self.assertEqual(self.client.patch(url, {"description": "x" * 1000}, format="json").status_code, 200)
        self.assertEqual(self.client.patch(url, {"description": "x" * 1001}, format="json").status_code, 400)

    def test_cleared_experience_date_can_be_saved_as_null(self):
        res = self.client.patch(f"{API}/about/{self.about.id}", {"experience_since": None}, format="json")
        self.assertEqual(res.status_code, 200, res.content)


class ExperienceYearsTests(APITestCase):
    """Years of experience are entered per portfolio and gain +1 every
    January 1 on their own."""

    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user("owner", password="pw")
        cls.profile = Profile.objects.create(user=cls.user, display_pic="https://img.example.com/p.png")

    def setUp(self):
        self.client = APIClient(HTTP_HOST="localhost")
        self.client.force_authenticate(user=self.user)

    def make_track(self, name, slug, years, entered_in=2026):
        track = PortfolioTrack.objects.create(profile=self.profile, name=name, slug=slug,
                                              display_pic="https://img.example.com/p.png")
        about = About.objects.create(track=track, skill=name, description="d",
                                     experience_base_years=years, experience_base_year=entered_in)
        return track, about

    @staticmethod
    def pretend_year(year):
        return mock.patch("portfolio.models.current_year", return_value=year)

    def api_years(self, about):
        return self.client.get(f"{API}/about/{about.id}").json()["years_of_experience"]

    def test_gains_one_year_every_january_first(self):
        _, about = self.make_track("Full Stack", "full-stack", years=4, entered_in=2026)
        for year, expected in [(2026, 4), (2027, 5), (2028, 6), (2036, 14)]:
            with self.pretend_year(year):
                self.assertEqual(self.api_years(about), expected, year)

    def test_same_year_does_not_change_and_never_goes_negative(self):
        _, about = self.make_track("Full Stack", "full-stack", years=4, entered_in=2026)
        with self.pretend_year(2026):
            self.assertEqual(about.years_of_experience, 4)
        with self.pretend_year(2020):
            self.assertEqual(about.years_of_experience, 4)

    def test_each_portfolio_has_its_own_years(self):
        _, a = self.make_track("Full Stack", "full-stack", years=4)
        _, b = self.make_track("CAD Engineer", "cad", years=1)
        with self.pretend_year(2027):
            self.assertEqual((self.api_years(a), self.api_years(b)), (5, 2))
        res = self.client.patch(f"{API}/about/{b.id}", {"experience_base_years": 3}, format="json")
        self.assertEqual(res.status_code, 200, res.content)
        a.refresh_from_db(); b.refresh_from_db()
        self.assertEqual((a.experience_base_years, b.experience_base_years), (4, 3))

    def test_saving_a_number_anchors_it_to_the_current_year(self):
        _, about = self.make_track("Full Stack", "full-stack", years=4, entered_in=2026)
        with self.pretend_year(2028):
            self.assertEqual(self.api_years(about), 6)          # what the form shows
            res = self.client.patch(f"{API}/about/{about.id}", {"experience_base_years": 6}, format="json")
            self.assertEqual(res.status_code, 200, res.content)
            self.assertEqual(res.json()["years_of_experience"], 6)
        about.refresh_from_db()
        self.assertEqual((about.experience_base_years, about.experience_base_year), (6, 2028))
        with self.pretend_year(2029):
            self.assertEqual(self.api_years(about), 7)

    def test_editing_other_fields_keeps_the_rollover_going(self):
        _, about = self.make_track("Full Stack", "full-stack", years=4, entered_in=2026)
        with self.pretend_year(2027):
            res = self.client.patch(f"{API}/about/{about.id}", {"description": "new text"}, format="json")
            self.assertEqual(res.status_code, 200, res.content)
        about.refresh_from_db()
        self.assertEqual((about.experience_base_years, about.experience_base_year), (4, 2026))
        with self.pretend_year(2027):
            self.assertEqual(self.api_years(about), 5)

    def test_typing_the_stored_number_again_after_new_year_sets_exactly_that(self):
        # Entered 2 in 2025; in 2026 the page shows 3. Typing 2 must show 2.
        _, about = self.make_track("CAD Engineer", "cad", years=2, entered_in=2025)
        with self.pretend_year(2026):
            self.assertEqual(self.api_years(about), 3)
            res = self.client.patch(f"{API}/about/{about.id}", {"experience_base_years": 2}, format="json")
            self.assertEqual(res.status_code, 200, res.content)
            self.assertEqual(res.json()["years_of_experience"], 2)
            self.assertEqual(self.api_years(about), 2)
        with self.pretend_year(2027):
            self.assertEqual(self.api_years(about), 3)

    def test_saving_the_displayed_number_unchanged_is_harmless(self):
        _, about = self.make_track("Full Stack", "full-stack", years=4, entered_in=2026)
        with self.pretend_year(2027):
            self.client.patch(f"{API}/about/{about.id}", {"experience_base_years": 5}, format="json")
            self.assertEqual(self.api_years(about), 5)
        with self.pretend_year(2028):
            self.assertEqual(self.api_years(about), 6)

    def test_clients_cannot_set_the_anchor_year(self):
        _, about = self.make_track("Full Stack", "full-stack", years=4, entered_in=2026)
        self.client.patch(f"{API}/about/{about.id}", {"experience_base_year": 1990}, format="json")
        about.refresh_from_db()
        self.assertEqual(about.experience_base_year, 2026)

    def test_years_validation(self):
        _, about = self.make_track("Full Stack", "full-stack", years=4)
        url = f"{API}/about/{about.id}"
        for bad in (-1, 81, 4.5, "abc", None):
            res = self.client.patch(url, {"experience_base_years": bad}, format="json")
            self.assertEqual(res.status_code, 400, bad)
        for good in (0, 80, 12):
            res = self.client.patch(url, {"experience_base_years": good}, format="json")
            self.assertEqual(res.status_code, 200, good)

    def test_duplicate_carries_years_and_then_stands_alone(self):
        source, about = self.make_track("Full Stack", "full-stack", years=4, entered_in=2026)
        with self.pretend_year(2027):
            res = self.client.post(f"{API}/tracks/", {"username": "owner", "name": "Copy", "copy_from": source.id},
                                   format="json")
            self.assertEqual(res.status_code, 201, res.content)
            copy_about = About.objects.get(track_id=res.json()["id"])
            self.assertEqual(self.api_years(copy_about), self.api_years(about))
            self.assertEqual(copy_about.years_of_experience, 5)

            self.client.patch(f"{API}/about/{copy_about.id}", {"experience_base_years": 9}, format="json")
            self.assertEqual(self.api_years(about), 5)
            self.client.patch(f"{API}/about/{about.id}", {"experience_base_years": 2}, format="json")
            copy_about.refresh_from_db()
            self.assertEqual(copy_about.years_of_experience, 9)

    def test_blank_track_starts_at_zero(self):
        res = self.client.post(f"{API}/tracks/", {"username": "owner", "name": "Blank"}, format="json")
        self.assertEqual(res.status_code, 201, res.content)
        self.assertEqual(res.json()["about"]["years_of_experience"], 0)

    def test_data_migration_converts_the_start_date(self):
        migration = importlib.import_module("portfolio.migrations.0014_experience_years_from_since")
        _, dated = self.make_track("Full Stack", "full-stack", years=0, entered_in=2000)
        dated.experience_since = date(2022, 7, 29)
        dated.save()
        _, undated = self.make_track("CAD", "cad", years=7, entered_in=2000)
        About.objects.filter(pk=undated.pk).update(experience_since=None)

        with mock.patch.object(migration, "date") as fake_date:
            fake_date.today.return_value = date(2026, 10, 18)
            migration.experience_years_from_start_date(django_apps, None)

        dated.refresh_from_db(); undated.refresh_from_db()
        self.assertEqual((dated.experience_base_years, dated.experience_base_year), (4, 2026))
        self.assertEqual((undated.experience_base_years, undated.experience_base_year), (0, 2026))


class EntryDescriptionTests(APITestCase):
    """Resume, education and skill entries carry an optional, capped
    description shown on the resume board."""

    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user("owner", password="pw")
        profile = Profile.objects.create(user=cls.user, display_pic="https://img.example.com/p.png")
        cls.track = PortfolioTrack.objects.create(profile=profile, name="Full Stack", slug="full-stack",
                                                  display_pic="https://img.example.com/p.png")
        cls.about = About.objects.create(track=cls.track, skill="s", description="d")
        cls.resume = Resume.objects.create(about=cls.about, post="Dev", company_name="Acme",
                                           start_year=date(2022, 1, 1), end_year=date(2023, 1, 1),
                                           description="Built things.")
        cls.school = School.objects.create(about=cls.about, school_name="Uni", start_year=date(2018, 1, 1),
                                           end_year=date(2022, 1, 1), description="Studied things.")
        cls.skill = Proficiency.objects.create(about=cls.about, skill_name="Python", skill_range=80,
                                               description="Used daily.")

    def setUp(self):
        self.client = APIClient(HTTP_HOST="localhost")
        self.client.force_authenticate(user=self.user)

    def urls(self):
        return {"resume": f"{API}/resume/{self.resume.id}", "school": f"{API}/school/{self.school.id}",
                "proficiency": f"{API}/proficiency/{self.skill.id}"}

    def test_cap_is_240(self):
        self.assertEqual(ENTRY_DESCRIPTION_MAX, 240)

    def test_description_is_optional_everywhere(self):
        for name, url in self.urls().items():
            res = self.client.patch(url, {"description": ""}, format="json")
            self.assertEqual(res.status_code, 200, (name, res.content))
        res = self.client.post(f"{API}/resume/", {"about": self.about.id, "post": "Dev 2", "company_name": "B",
                                                  "start_year": "2023-01-01", "end_year": "2024-01-01"}, format="json")
        self.assertEqual(res.status_code, 201, res.content)
        self.assertEqual(res.json()["description"], "")

    def test_description_cap_enforced_on_all_three(self):
        for name, url in self.urls().items():
            ok = self.client.patch(url, {"description": "x" * 240}, format="json")
            self.assertEqual(ok.status_code, 200, (name, ok.content))
            bad = self.client.patch(url, {"description": "x" * 241}, format="json")
            self.assertEqual(bad.status_code, 400, name)
            self.assertIn("description", bad.json())

    def test_descriptions_are_in_the_public_payload(self):
        data = self.client.get(f"{API}/profiles/?username=owner").json()[0]
        about = data["tracks"][0]["about"]
        self.assertEqual(about["resume"][0]["description"], "Built things.")
        self.assertEqual(about["school"][0]["description"], "Studied things.")
        self.assertEqual(about["proficiency"][0]["description"], "Used daily.")

    def test_duplicate_copies_descriptions_and_they_then_diverge(self):
        res = self.client.post(f"{API}/tracks/", {"username": "owner", "name": "Copy", "copy_from": self.track.id},
                               format="json")
        self.assertEqual(res.status_code, 201, res.content)
        copy_about = About.objects.get(track_id=res.json()["id"])
        self.assertEqual(copy_about.resume.get().description, "Built things.")
        self.assertEqual(copy_about.school.get().description, "Studied things.")
        self.assertEqual(copy_about.proficiency.get().description, "Used daily.")

        self.client.patch(f"{API}/resume/{copy_about.resume.get().id}", {"description": "Copy only"}, format="json")
        self.resume.refresh_from_db()
        self.assertEqual(self.resume.description, "Built things.")


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])   # fast, tests only
class ChangePasswordTests(APITestCase):
    OLD = "Old-Passw0rd-xyz"
    NEW = "Brand-New-Passw0rd-42"

    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user("owner", password=cls.OLD, first_name="Own", last_name="Er")
        Profile.objects.create(user=cls.user, display_pic="https://img.example.com/p.png")

    def setUp(self):
        cache.clear()      # the endpoint is rate limited; keep tests independent
        self.url = f"{API}/auth/change-password"

    def session(self, password=None):
        client = APIClient(HTTP_HOST="localhost")
        res = client.post(f"{API}/auth/login", {"username": "owner", "password": password or self.OLD}, format="json")
        self.assertEqual(res.status_code, 200, res.content)
        return client

    def me(self, client):
        return client.get(f"{API}/auth/me").json()["authenticated"]

    def change(self, client, current, new):
        return client.post(self.url, {"current_password": current, "new_password": new}, format="json")

    def test_wrong_current_password_is_rejected_and_nothing_changes(self):
        client = self.session()
        res = self.change(client, "not-the-password", self.NEW)
        self.assertEqual(res.status_code, 400)
        self.assertIn("current_password", res.json())
        self.user.refresh_from_db()
        self.assertTrue(self.user.check_password(self.OLD))

    def test_weak_or_unchanged_new_passwords_are_rejected(self):
        client = self.session()
        for weak in ("short1", "password123", "12345678901", "owner1234", self.OLD):
            res = self.change(client, self.OLD, weak)
            self.assertEqual(res.status_code, 400, weak)
            self.assertIn("new_password", res.json(), weak)
        self.user.refresh_from_db()
        self.assertTrue(self.user.check_password(self.OLD))

    def test_missing_fields_are_rejected(self):
        client = self.session()
        res = client.post(self.url, {}, format="json")
        self.assertEqual(res.status_code, 400)
        self.assertEqual(set(res.json()), {"current_password", "new_password"})

    def test_change_works_and_the_same_session_keeps_going(self):
        client = self.session()
        res = self.change(client, self.OLD, self.NEW)
        self.assertEqual(res.status_code, 200, res.content)
        self.assertIn("access_token", res.cookies)
        self.assertIn("refresh_token", res.cookies)
        self.assertTrue(self.me(client), "current session should stay signed in")
        # and it can still write
        track = PortfolioTrack.objects.create(profile=Profile.objects.get(), name="T", slug="t",
                                              display_pic="https://img.example.com/p.png")
        res = client.patch(f"{API}/tracks/{track.id}", {"name": "T2"}, format="json")
        self.assertEqual(res.status_code, 200, res.content)

    def test_old_password_stops_working_and_new_one_logs_in(self):
        self.change(self.session(), self.OLD, self.NEW)
        anon = APIClient(HTTP_HOST="localhost")
        self.assertEqual(anon.post(f"{API}/auth/login", {"username": "owner", "password": self.OLD},
                                   format="json").status_code, 401)
        self.assertEqual(anon.post(f"{API}/auth/login", {"username": "owner", "password": self.NEW},
                                   format="json").status_code, 200)

    def test_other_sessions_are_signed_out_but_can_log_in_again(self):
        phone, laptop = self.session(), self.session()
        self.assertTrue(self.me(phone) and self.me(laptop))

        self.assertEqual(self.change(laptop, self.OLD, self.NEW).status_code, 200)

        self.assertFalse(self.me(phone), "old session must be revoked")
        track = PortfolioTrack.objects.create(profile=Profile.objects.get(), name="T", slug="t",
                                              display_pic="https://img.example.com/p.png")
        self.assertIn(phone.patch(f"{API}/tracks/{track.id}", {"name": "Hacked"}, format="json").status_code,
                      (401, 403))
        # public reads and logging in again still work with the stale cookies attached
        self.assertEqual(phone.get(f"{API}/tracks/").status_code, 200)
        relog = phone.post(f"{API}/auth/login", {"username": "owner", "password": self.NEW}, format="json")
        self.assertEqual(relog.status_code, 200, relog.content)
        self.assertTrue(self.me(phone))

    def test_anonymous_cannot_change_a_password(self):
        res = self.change(APIClient(HTTP_HOST="localhost"), self.OLD, self.NEW)
        self.assertIn(res.status_code, (401, 403))
        self.user.refresh_from_db()
        self.assertTrue(self.user.check_password(self.OLD))

    def test_guessing_is_rate_limited(self):
        client = self.session()
        codes = [self.change(client, f"guess-{i}", self.NEW).status_code for i in range(7)]
        self.assertEqual(codes[:5], [400] * 5)
        self.assertEqual(codes[5], 429)

    def test_password_hash_is_never_in_the_public_profile(self):
        anon = APIClient(HTTP_HOST="localhost")
        for url in (f"{API}/profiles/?username=owner", f"{API}/profiles/{Profile.objects.get().id}"):
            body = anon.get(url).content.decode()
            self.assertNotIn("pbkdf2", body, url)
            self.assertNotIn('"password"', body, url)


class SeedEntryDescriptionsTests(APITestCase):
    """The one-off command that drafts descriptions for existing entries."""

    @classmethod
    def setUpTestData(cls):
        user = User.objects.create_user("owner", password="pw")
        profile = Profile.objects.create(user=user, display_pic="https://img.example.com/p.png")
        track = PortfolioTrack.objects.create(profile=profile, name="Full Stack", slug="full-stack",
                                              display_pic="https://img.example.com/p.png")
        about = About.objects.create(track=track, skill="s", description="d")
        cls.job = Resume.objects.create(about=about, post="Robotic Instructor", company_name="Al Izzah",
                                        start_year=date(2022, 9, 1), end_year=date(2022, 10, 1))
        cls.written = Resume.objects.create(about=about, post="Class Teacher", company_name="Al Izzah",
                                            start_year=date(2022, 10, 1), end_year=date(2023, 1, 1),
                                            description="My own words.")
        cls.unknown = Resume.objects.create(about=about, post="Astronaut", company_name="NASA",
                                            start_year=date(2020, 1, 1), end_year=date(2021, 1, 1))
        cls.skill = Proficiency.objects.create(about=about, skill_name="Python", skill_range=80)
        cls.school = School.objects.create(about=about, school_name="Uni", start_year=date(2018, 1, 1),
                                           end_year=date(2022, 1, 1))

    def run_command(self, *args):
        out = StringIO()
        call_command("seed_entry_descriptions", *args, stdout=out)
        return out.getvalue()

    def test_dry_run_changes_nothing(self):
        output = self.run_command()
        self.assertIn("dry run", output)
        self.job.refresh_from_db()
        self.assertEqual(self.job.description, "")

    def test_apply_fills_blanks_only(self):
        self.run_command("--apply")
        for entry in (self.job, self.written, self.unknown, self.skill, self.school):
            entry.refresh_from_db()
        self.assertTrue(self.job.description)
        self.assertTrue(self.skill.description)
        self.assertEqual(self.written.description, "My own words.", "never overwrites what the owner wrote")
        self.assertEqual(self.unknown.description, "", "entries without a draft are left alone")
        self.assertEqual(self.school.description, "", "education is left for the owner to write")

    def test_is_idempotent(self):
        self.run_command("--apply")
        self.assertIn("0 entries", self.run_command("--apply"))

    def test_every_draft_fits_the_cap(self):
        from portfolio.management.commands import seed_entry_descriptions as seed
        for text in list(seed.RESUME_DESCRIPTIONS.values()) + list(seed.SKILL_DESCRIPTIONS.values()):
            self.assertLessEqual(len(text), ENTRY_DESCRIPTION_MAX, text)


BROWSER = "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36"


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])   # fast, tests only
class VisitorAnalyticsTests(APITestCase):
    """Anonymous visitor counting: the owner and bots don't count, a returning
    person is a repeat on the same visitor, and only the owner can read it."""

    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user("owner", password="Old-Passw0rd-xyz")
        cls.profile = Profile.objects.create(user=cls.user, display_pic="https://img.example.com/p.png")
        cls.fs = PortfolioTrack.objects.create(profile=cls.profile, name="Full Stack", slug="full-stack",
                                               display_pic="https://img.example.com/p.png")
        cls.cad = PortfolioTrack.objects.create(profile=cls.profile, name="CAD Engineer", slug="cad",
                                                display_pic="https://img.example.com/p.png")

    def setUp(self):
        cache.clear()
        self.url = f"{API}/analytics/visit"
        self.summary = f"{API}/analytics/summary"

    def browser(self, agent=BROWSER):
        return APIClient(HTTP_HOST="localhost", HTTP_USER_AGENT=agent)

    def visit(self, vid, path="/", client=None):
        return (client or self.browser()).post(self.url, {"visitor_id": str(vid), "path": path}, format="json")

    def owner_client(self):
        client = self.browser()
        res = client.post(f"{API}/auth/login", {"username": "owner", "password": "Old-Passw0rd-xyz"}, format="json")
        self.assertEqual(res.status_code, 200, res.content)
        return client

    def at(self, moment):
        return mock.patch("portfolio.analytics.timezone.now", return_value=moment)

    # --- counting ------------------------------------------------------------

    def test_a_first_visit_creates_one_visitor_and_one_page_view(self):
        vid = uuid.uuid4()
        self.assertEqual(self.visit(vid).status_code, 204)
        visitor = Visitor.objects.get()
        self.assertEqual((str(visitor.visitor_id), visitor.visit_count), (str(vid), 1))
        view = PageView.objects.get()
        self.assertIsNone(view.track)

    def test_repeating_the_same_page_in_one_visit_changes_nothing(self):
        vid = uuid.uuid4()
        for _ in range(3):
            self.visit(vid)
            self.visit(vid, "full-stack")
        self.assertEqual((Visitor.objects.count(), Visitor.objects.get().visit_count), (1, 1))
        self.assertEqual(PageView.objects.count(), 2)       # landing + full-stack, once each

    def test_browsing_more_portfolios_in_the_same_visit_is_not_a_new_visit(self):
        vid = uuid.uuid4()
        self.visit(vid); self.visit(vid, "full-stack"); self.visit(vid, "/cad")
        visitor = Visitor.objects.get()
        self.assertEqual(visitor.visit_count, 1)
        self.assertEqual(set(visitor.views.values_list("track__slug", flat=True)), {None, "full-stack", "cad"})

    def test_coming_back_later_is_a_repeat_on_the_same_visitor_not_a_new_one(self):
        vid, start = uuid.uuid4(), timezone.now()
        with self.at(start):
            self.visit(vid, "full-stack")
        with self.at(start + timedelta(minutes=10)):
            self.visit(vid, "cad")                          # same session
        self.assertEqual(Visitor.objects.get().visit_count, 1)
        with self.at(start + timedelta(minutes=10 + 31)):
            self.visit(vid, "full-stack")                   # a fresh visit, same person
        self.assertEqual(Visitor.objects.count(), 1, "unique visitors must not go up")
        self.assertEqual(Visitor.objects.get().visit_count, 2)
        self.assertEqual(PageView.objects.filter(track=self.fs).count(), 2, "the repeat is recorded")

    def test_different_browsers_are_different_visitors(self):
        self.visit(uuid.uuid4()); self.visit(uuid.uuid4())
        self.assertEqual(Visitor.objects.count(), 2)

    # --- who is NOT counted ---------------------------------------------------

    def test_the_signed_in_owner_is_never_counted(self):
        owner = self.owner_client()
        res = self.visit(uuid.uuid4(), "full-stack", client=owner)
        self.assertEqual(res.status_code, 204)
        self.assertEqual((Visitor.objects.count(), PageView.objects.count()), (0, 0))

    def test_an_owner_whose_session_was_revoked_is_just_a_visitor(self):
        owner = self.owner_client()
        stale = {c: owner.cookies[c].value for c in ("access_token", "refresh_token")}
        self.user.set_password("Some-Other-Pass-9"); self.user.save()      # revokes the old tokens
        client = self.browser()
        for name, value in stale.items():
            client.cookies[name] = value
        self.visit(uuid.uuid4(), client=client)
        self.assertEqual(Visitor.objects.count(), 1)

    def test_bots_and_automation_are_not_counted(self):
        for agent in ("Mozilla/5.0 (compatible; Googlebot/2.1; +http://www.google.com/bot.html)",
                      "Mozilla/5.0 HeadlessChrome/120.0", "Slackbot-LinkExpanding 1.0", "python-requests/2.31",
                      "facebookexternalhit/1.1", "Chrome-Lighthouse", ""):
            self.assertEqual(self.visit(uuid.uuid4(), client=self.browser(agent)).status_code, 204, agent)
        self.assertEqual(Visitor.objects.count(), 0)
        self.visit(uuid.uuid4())
        self.assertEqual(Visitor.objects.count(), 1)

    def test_unknown_pages_are_ignored_and_look_identical_to_counted_ones(self):
        res = self.visit(uuid.uuid4(), "no-such-portfolio")
        self.assertEqual(res.status_code, 204)
        self.assertEqual(Visitor.objects.count(), 0)

    def test_bad_requests(self):
        self.assertEqual(self.browser().post(self.url, {"path": "/"}, format="json").status_code, 400)
        self.assertEqual(self.visit("not-a-uuid").status_code, 400)
        res = self.browser().post(self.url, {"visitor_id": str(uuid.uuid4()), "path": {"x": 1}}, format="json")
        self.assertEqual(res.status_code, 204)
        self.assertEqual(Visitor.objects.count(), 0)

    def test_the_beacon_is_rate_limited_per_ip(self):
        with mock.patch.dict(ScopedRateThrottle.THROTTLE_RATES, {"visit": "3/hour"}):
            codes = [self.visit(uuid.uuid4()).status_code for _ in range(5)]
        self.assertEqual(codes, [204, 204, 204, 429, 429])
        self.assertEqual(Visitor.objects.count(), 3)

    # --- who can SEE it ---------------------------------------------------------

    def test_only_the_signed_in_owner_can_read_the_summary(self):
        self.assertIn(self.browser().get(self.summary).status_code, (401, 403))
        self.assertEqual(self.owner_client().get(self.summary).status_code, 200)

    def test_summary_numbers(self):
        a, b, c, t0 = uuid.uuid4(), uuid.uuid4(), uuid.uuid4(), timezone.now() - timedelta(hours=3)
        with self.at(t0):
            self.visit(a); self.visit(a, "full-stack")      # A: visit 1
            self.visit(b, "cad")
            self.visit(c)
        with self.at(t0 + timedelta(hours=2)):
            self.visit(a, "cad")                            # A comes back: visit 2
        data = self.owner_client().get(self.summary).json()

        self.assertEqual(data["unique_visitors"], 3, "A counted once even though they came twice")
        self.assertEqual(data["returning_visitors"], 1)
        self.assertEqual(data["total_visits"], 4)
        self.assertEqual(data["new_last_7_days"], 3)
        self.assertEqual(data["landing"], {"viewers": 2, "visits": 2})
        by_slug = {p["slug"]: p for p in data["portfolios"]}
        self.assertEqual(by_slug["full-stack"], {"slug": "full-stack", "name": "Full Stack", "viewers": 1, "visits": 1})
        self.assertEqual((by_slug["cad"]["viewers"], by_slug["cad"]["visits"]), (2, 2))

        recent = data["recent"]
        self.assertEqual(len(recent), 3)
        first = recent[0]                                    # A was seen most recently
        self.assertTrue(first["returning"] and first["visits"] == 2)
        self.assertEqual({v["slug"]: v["visits"] for v in first["viewed"]}, {"full-stack": 1, "cad": 1})
        self.assertEqual(first["landing_visits"], 1)
        self.assertFalse(any(r["returning"] for r in recent[1:]))

    def test_summary_exposes_no_browser_ids_or_network_details(self):
        vid = uuid.uuid4()
        self.visit(vid)
        body = self.owner_client().get(self.summary).content.decode()
        self.assertNotIn(str(vid), body)
        self.assertEqual({f.name for f in Visitor._meta.get_fields()} & {"ip", "ip_address", "user_agent"}, set())

    def test_deleting_a_portfolio_keeps_the_visitors(self):
        vid = uuid.uuid4()
        self.visit(vid, "full-stack")
        self.fs.delete()
        self.assertEqual((Visitor.objects.count(), PageView.objects.count()), (1, 0))

    # --- resume views -------------------------------------------------------------------

    def open_resume(self, vid, slug="full-stack", client=None):
        return (client or self.browser()).post(f"{API}/analytics/resume-view", {"visitor_id": str(vid), "slug": slug}, format="json")

    def test_opening_a_resume_is_counted_per_portfolio_with_the_time(self):
        vid, t0 = uuid.uuid4(), timezone.now() - timedelta(hours=5)
        with self.at(t0):
            self.assertEqual(self.open_resume(vid, "full-stack").status_code, 204)
            self.assertEqual(self.open_resume(vid, "cad").status_code, 204)
        row = ResumeView.objects.get(track=self.fs)
        self.assertEqual((row.visitor.visitor_id, row.created_at), (vid, t0))
        self.assertEqual(ResumeView.objects.count(), 2)
        self.assertEqual(Visitor.objects.count(), 1, "the same browser is one visitor")

    def test_a_refresh_is_not_counted_again_but_coming_back_later_is(self):
        vid, t0 = uuid.uuid4(), timezone.now() - timedelta(hours=5)
        with self.at(t0):
            self.open_resume(vid)
        with self.at(t0 + timedelta(minutes=10)):
            self.open_resume(vid)
        self.assertEqual(ResumeView.objects.count(), 1)
        with self.at(t0 + timedelta(hours=3)):
            self.open_resume(vid)
        self.assertEqual(ResumeView.objects.count(), 2)

    def test_two_people_opening_the_same_resume_are_two_views(self):
        self.open_resume(uuid.uuid4()); self.open_resume(uuid.uuid4())
        self.assertEqual((ResumeView.objects.count(), Visitor.objects.count()), (2, 2))

    def test_the_owner_bots_and_unknown_portfolios_are_not_counted(self):
        self.assertEqual(self.open_resume(uuid.uuid4(), client=self.owner_client()).status_code, 204)
        for agent in ("Googlebot/2.1", "HeadlessChrome/120.0", "python-requests/2.31", ""):
            self.assertEqual(self.open_resume(uuid.uuid4(), client=self.browser(agent)).status_code, 204, agent)
        self.assertEqual(self.open_resume(uuid.uuid4(), "no-such-portfolio").status_code, 204)
        self.assertEqual(self.browser().post(f"{API}/analytics/resume-view", {"visitor_id": str(uuid.uuid4()), "slug": {"x": 1}}, format="json").status_code, 204)
        self.assertEqual((ResumeView.objects.count(), Visitor.objects.count()), (0, 0))

    def test_a_bad_visitor_id_is_a_400(self):
        self.assertEqual(self.browser().post(f"{API}/analytics/resume-view", {"slug": "full-stack"}, format="json").status_code, 400)
        self.assertEqual(self.open_resume("not-a-uuid").status_code, 400)
        self.assertEqual(ResumeView.objects.count(), 0)

    def test_the_resume_beacon_is_rate_limited(self):
        with mock.patch.dict(ScopedRateThrottle.THROTTLE_RATES, {"visit": "3/hour"}):
            codes = [self.open_resume(uuid.uuid4()).status_code for _ in range(5)]
        self.assertEqual(codes, [204, 204, 204, 429, 429])

    def test_the_summary_shows_how_many_times_which_resume_and_when(self):
        a, b, t0 = uuid.uuid4(), uuid.uuid4(), timezone.now() - timedelta(hours=6)
        with self.at(t0):
            self.open_resume(a, "full-stack")
        with self.at(t0 + timedelta(hours=1)):
            self.open_resume(b, "full-stack")
        with self.at(t0 + timedelta(hours=2)):
            self.open_resume(a, "cad")
        with self.at(t0 + timedelta(hours=4)):
            self.open_resume(a, "full-stack")                 # A comes back to the first one
        data = self.owner_client().get(self.summary).json()["resume"]

        self.assertEqual((data["total_views"], data["viewers"]), (4, 2))
        by_slug = {row["slug"]: row for row in data["portfolios"]}
        self.assertEqual((by_slug["full-stack"]["views"], by_slug["full-stack"]["viewers"]), (3, 2))
        self.assertEqual((by_slug["cad"]["views"], by_slug["cad"]["viewers"]), (1, 1))
        self.assertEqual(by_slug["full-stack"]["last_viewed"], (t0 + timedelta(hours=4)).isoformat().replace("+00:00", "Z"))
        recent = data["recent"]
        self.assertEqual([r["slug"] for r in recent], ["full-stack", "cad", "full-stack", "full-stack"], "newest first")
        self.assertEqual(recent[0]["name"], "Full Stack")
        self.assertTrue(all(r["at"] for r in recent))

        people = {tuple(sorted((v["slug"], v["views"]) for v in person["resume_views"])) for person in self.owner_client().get(self.summary).json()["recent"]}
        self.assertEqual(people, {(("cad", 1), ("full-stack", 2)), (("full-stack", 1),)})

    def test_an_unviewed_portfolio_reads_zero_not_missing(self):
        data = self.owner_client().get(self.summary).json()["resume"]
        self.assertEqual((data["total_views"], data["viewers"], data["recent"]), (0, 0, []))
        self.assertEqual([(r["slug"], r["views"], r["last_viewed"]) for r in data["portfolios"]], [("full-stack", 0, None), ("cad", 0, None)])

    def test_the_summary_lists_only_the_latest_resume_views(self):
        t0 = timezone.now() - timedelta(days=1)
        for i in range(20):
            with self.at(t0 + timedelta(minutes=i)):
                self.open_resume(uuid.uuid4())
        data = self.owner_client().get(self.summary).json()["resume"]
        self.assertEqual((data["total_views"], len(data["recent"])), (20, 15))

    def test_resume_views_hold_nothing_personal_and_follow_their_portfolio(self):
        vid = uuid.uuid4()
        self.open_resume(vid)
        self.assertNotIn(str(vid), self.owner_client().get(self.summary).content.decode())
        self.assertEqual({f.name for f in ResumeView._meta.get_fields()} & {"ip", "ip_address", "user_agent"}, set())
        self.fs.delete()
        self.assertEqual((ResumeView.objects.count(), Visitor.objects.count()), (0, 1))

    def test_only_the_owner_sees_resume_views(self):
        self.open_resume(uuid.uuid4())
        self.assertIn(self.browser().get(self.summary).status_code, (401, 403))


class ProjectSharingTests(APITestCase):
    """One project can be shown in several portfolios: edited once, removed per portfolio,
    and deleted for good only when no portfolio shows it any more."""

    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user("owner", password="Old-Passw0rd-xyz")
        cls.profile = Profile.objects.create(user=cls.user, display_pic="https://img.example.com/p.png")
        cls.a, cls.b, cls.c = (PortfolioTrack.objects.create(profile=cls.profile, name=n, slug=s, display_pic="https://img.example.com/p.png")
                               for n, s in (("Alpha", "alpha"), ("Beta", "beta"), ("Gamma", "gamma")))
        for track in (cls.a, cls.b, cls.c):
            About.objects.create(track=track, skill="Skill", description="About")
        other = User.objects.create_user("stranger", password="Old-Passw0rd-xyz")
        cls.stranger_track = PortfolioTrack.objects.create(profile=Profile.objects.create(user=other, display_pic="https://img.example.com/s.png"),
                                                           name="Theirs", slug="theirs", display_pic="https://img.example.com/s.png")

    def setUp(self):
        self.client = APIClient(HTTP_HOST="localhost")
        self.client.force_authenticate(user=self.user)

    def body(self, **extra):
        return {"track": self.a.id, "category": "WebDev", "name": "Demo", "thumbnail": "https://a.com/t.png",
                "project_link": "https://demo.example.com", **extra}

    def add(self, **extra):
        res = self.client.post(f"{API}/portfolio/", self.body(**extra), format="json")
        self.assertEqual(res.status_code, 201, res.content)
        return Portfolio.objects.get(pk=res.json()["id"])

    def shown_in(self, project):
        return sorted(project.tracks.values_list("name", flat=True))

    def listed(self, track):
        tracks = APIClient(HTTP_HOST="localhost").get(f"{API}/tracks/?username=owner").json()
        return sorted(i["name"] for i in next(t for t in tracks if t["id"] == track.id)["items"])

    # --- adding -----------------------------------------------------------------------------

    def test_a_new_project_is_shown_only_where_it_was_posted_by_default(self):
        project = self.add()
        self.assertEqual(self.shown_in(project), ["Alpha"])
        self.assertEqual((self.listed(self.a), self.listed(self.b), self.listed(self.c)), (["Demo"], [], []))

    def test_it_can_be_shared_with_chosen_portfolios(self):
        project = self.add(share_with=[self.c.id])
        self.assertEqual(self.shown_in(project), ["Alpha", "Gamma"])
        self.assertEqual((self.listed(self.a), self.listed(self.b), self.listed(self.c)), (["Demo"], [], ["Demo"]))

    def test_share_with_all_means_every_portfolio_of_the_owner(self):
        project = self.add(share_with="all")
        self.assertEqual(self.shown_in(project), ["Alpha", "Beta", "Gamma"])
        self.assertEqual(self.stranger_track.items.count(), 0, "never another owner's portfolio")

    def test_it_is_one_record_whichever_portfolios_show_it(self):
        self.add(share_with=[self.b.id, self.c.id])
        self.assertEqual(Portfolio.objects.count(), 1)

    def test_sharing_is_validated(self):
        for bad, field in (([9999], "share_with"), ("everything", "share_with"), ([True], "share_with"), (["1"], "share_with"),
                           ([self.stranger_track.id], "tracks")):
            res = self.client.post(f"{API}/portfolio/", self.body(share_with=bad), format="json")
            self.assertEqual(res.status_code, 400, bad)
            self.assertIn(field, res.json(), bad)
        self.assertEqual(Portfolio.objects.count(), 0, "nothing half-created")
        res = self.client.post(f"{API}/portfolio/", {**self.body(), "track": None}, format="json")
        self.assertEqual(res.status_code, 400)
        res = self.client.post(f"{API}/portfolio/", {k: v for k, v in self.body().items() if k != "track"}, format="json")
        self.assertEqual(res.status_code, 400, "it has to say where it goes")
        self.assertEqual(self.client.post(f"{API}/portfolio/", self.body(track=self.stranger_track.id, share_with=[self.a.id]), format="json").status_code, 400)

    def test_the_old_way_of_adding_still_works_and_listing_the_portfolios_directly_does_too(self):
        res = self.client.post(f"{API}/portfolio/", {k: v for k, v in self.body().items() if k != "track"} | {"tracks": [self.b.id, self.c.id]}, format="json")
        self.assertEqual(res.status_code, 201, res.content)
        self.assertEqual(sorted(res.json()["tracks"]), sorted([self.b.id, self.c.id]))

    def test_the_payload_says_where_each_project_is_shown(self):
        project = self.add(share_with=[self.b.id])
        tracks = APIClient(HTTP_HOST="localhost").get(f"{API}/tracks/?username=owner").json()
        item = next(i for t in tracks if t["id"] == self.b.id for i in t["items"])
        self.assertEqual(sorted(item["tracks"]), sorted([self.a.id, self.b.id]))
        self.assertEqual(item["display_username"], "owner")
        self.assertEqual(item["id"], project.id)

    # --- editing ---------------------------------------------------------------------------------

    def test_editing_a_shared_project_changes_it_everywhere(self):
        project = self.add(share_with="all")
        self.client.patch(f"{API}/portfolio/{project.id}", {"name": "Renamed", "description": "New words."}, format="json")
        self.assertEqual((self.listed(self.a), self.listed(self.b), self.listed(self.c)), (["Renamed"],) * 3)
        project.refresh_from_db()
        self.assertEqual(project.description, "New words.")
        self.assertEqual(self.shown_in(project), ["Alpha", "Beta", "Gamma"], "editing the words leaves the sharing alone")

    def test_a_project_added_to_one_portfolio_can_be_added_to_another_later_without_retyping(self):
        project = self.add()
        res = self.client.patch(f"{API}/portfolio/{project.id}", {"tracks": [self.a.id, self.b.id]}, format="json")
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(sorted(res.json()["tracks"]), sorted([self.a.id, self.b.id]))
        self.assertEqual(self.listed(self.b), ["Demo"])
        self.assertEqual(Portfolio.objects.count(), 1, "the same project, not a copy")

    def test_unticking_one_portfolio_leaves_the_others_showing_it(self):
        project = self.add(share_with=[self.b.id, self.c.id])
        res = self.client.patch(f"{API}/portfolio/{project.id}", {"tracks": [self.a.id, self.c.id]}, format="json")
        self.assertEqual(res.status_code, 200)
        self.assertEqual((self.listed(self.a), self.listed(self.b), self.listed(self.c)), (["Demo"], [], ["Demo"]))
        self.assertTrue(Portfolio.objects.filter(pk=project.pk).exists())

    def test_unticking_every_portfolio_deletes_the_project(self):
        project = self.add(share_with=[self.b.id])
        self.client.patch(f"{API}/portfolio/{project.id}", {"tracks": [self.b.id]}, format="json")
        self.assertTrue(Portfolio.objects.filter(pk=project.pk).exists(), "still shown in Beta")
        res = self.client.patch(f"{API}/portfolio/{project.id}", {"tracks": []}, format="json")
        self.assertEqual(res.status_code, 204)
        self.assertFalse(Portfolio.objects.filter(pk=project.pk).exists())
        self.assertEqual(self.client.get(f"{API}/portfolio/{project.id}").status_code, 404)

    def test_swapping_portfolios_in_one_save_never_loses_the_project(self):
        project = self.add()
        res = self.client.patch(f"{API}/portfolio/{project.id}", {"tracks": [self.b.id]}, format="json")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(self.shown_in(Portfolio.objects.get(pk=project.pk)), ["Beta"])

    def test_sharing_with_another_owner_is_refused_on_edit(self):
        project = self.add()
        res = self.client.patch(f"{API}/portfolio/{project.id}", {"tracks": [self.a.id, self.stranger_track.id]}, format="json")
        self.assertEqual(res.status_code, 400)
        self.assertEqual(self.shown_in(project), ["Alpha"])
        self.assertEqual(self.client.patch(f"{API}/portfolio/{project.id}", {"tracks": [9999]}, format="json").status_code, 400)

    def test_anonymous_visitors_cannot_change_who_shows_a_project(self):
        project = self.add()
        anon = APIClient(HTTP_HOST="localhost")
        self.assertIn(anon.patch(f"{API}/portfolio/{project.id}", {"tracks": [self.b.id]}, format="json").status_code, (401, 403))
        self.assertIn(anon.delete(f"{API}/portfolio/{project.id}?from={self.a.id}").status_code, (401, 403))
        self.assertIn(anon.post(f"{API}/portfolio/", self.body(share_with="all"), format="json").status_code, (401, 403))
        self.assertEqual(self.shown_in(project), ["Alpha"])

    # --- deleting --------------------------------------------------------------------------------

    def delete(self, project, **query):
        qs = "&".join(f"{k}={v}" for k, v in query.items())
        return self.client.delete(f"{API}/portfolio/{project.id}" + (f"?{qs}" if qs else ""))

    def test_deleting_from_one_portfolio_leaves_the_others(self):
        project = self.add(share_with="all")
        self.assertEqual(self.delete(project, **{"from": self.a.id}).status_code, 204)
        self.assertEqual((self.listed(self.a), self.listed(self.b), self.listed(self.c)), ([], ["Demo"], ["Demo"]))
        self.assertTrue(Portfolio.objects.filter(pk=project.pk).exists())

    def test_deleting_can_also_remove_it_from_chosen_portfolios(self):
        project = self.add(share_with="all")
        self.assertEqual(self.delete(project, **{"from": self.a.id, "also": self.c.id}).status_code, 204)
        self.assertEqual((self.listed(self.a), self.listed(self.b), self.listed(self.c)), ([], ["Demo"], []))

    def test_deleting_from_every_portfolio_deletes_it_for_good(self):
        project = self.add(share_with="all")
        self.assertEqual(self.delete(project, **{"from": self.b.id, "also": "all"}).status_code, 204)
        self.assertFalse(Portfolio.objects.filter(pk=project.pk).exists())
        again = self.add(share_with=[self.b.id])
        self.delete(again, **{"from": self.a.id})
        self.assertTrue(Portfolio.objects.filter(pk=again.pk).exists())
        self.delete(again, **{"from": self.b.id})
        self.assertFalse(Portfolio.objects.filter(pk=again.pk).exists(), "removed from the last portfolio: gone")

    def test_a_shared_project_cannot_be_deleted_without_saying_from_where(self):
        project = self.add(share_with="all")
        res = self.delete(project)
        self.assertEqual(res.status_code, 400)
        self.assertIn("from", res.json())
        self.assertEqual(self.shown_in(project), ["Alpha", "Beta", "Gamma"], "nothing was removed")

    def test_a_project_in_one_portfolio_can_be_deleted_either_way(self):
        one, two = self.add(), self.add(name="Other")
        self.assertEqual(self.delete(one).status_code, 204)
        self.assertEqual(self.delete(two, **{"from": self.a.id}).status_code, 204)
        self.assertEqual(Portfolio.objects.count(), 0)

    def test_bad_delete_requests_change_nothing(self):
        project = self.add(share_with=[self.b.id])
        for query in ({"from": self.c.id}, {"from": self.a.id, "also": self.c.id}, {"from": "x"}, {"from": self.a.id, "also": "1,x"}):
            self.assertEqual(self.delete(project, **query).status_code, 400, query)
        self.assertEqual(self.shown_in(project), ["Alpha", "Beta"])

    # --- deleting a whole portfolio -----------------------------------------------------------------

    def test_deleting_a_portfolio_keeps_the_projects_other_portfolios_show(self):
        shared, only_alpha = self.add(share_with=[self.b.id]), self.add(name="Only Alpha")
        self.assertEqual(self.client.delete(f"{API}/tracks/{self.a.id}").status_code, 204)
        self.assertTrue(Portfolio.objects.filter(pk=shared.pk).exists())
        self.assertFalse(Portfolio.objects.filter(pk=only_alpha.pk).exists(), "it was shown nowhere else")
        self.assertEqual(self.listed(self.b), ["Demo"])
        self.assertEqual(Portfolio.objects.count(), 1)

    # --- copying a whole portfolio ------------------------------------------------------------------

    def test_duplicating_a_portfolio_copies_the_projects_so_they_stand_alone(self):
        project = self.add(share_with=[self.b.id])
        res = self.client.post(f"{API}/tracks/", {"username": "owner", "name": "Copy", "copy_from": self.a.id, "copy_portfolio_items": True}, format="json")
        self.assertEqual(res.status_code, 201, res.content)
        copy = PortfolioTrack.objects.get(pk=res.json()["id"])
        self.assertEqual(copy.items.count(), 1)
        self.assertEqual(Portfolio.objects.count(), 2)
        self.client.patch(f"{API}/portfolio/{copy.items.get().id}", {"name": "Changed in the copy"}, format="json")
        project.refresh_from_db()
        self.assertEqual(project.name, "Demo")

    def test_a_new_portfolio_can_share_the_originals_projects_instead_of_copying_them(self):
        project = self.add(share_with=[self.b.id])
        res = self.client.post(f"{API}/tracks/", {"username": "owner", "name": "Linked", "copy_from": self.a.id, "share_portfolio_items": True}, format="json")
        self.assertEqual(res.status_code, 201, res.content)
        linked = PortfolioTrack.objects.get(pk=res.json()["id"])
        self.assertEqual(Portfolio.objects.count(), 1, "no new project rows: the same ones")
        self.assertEqual(self.shown_in(project), ["Alpha", "Beta", "Linked"])
        self.assertEqual([i["name"] for i in res.json()["items"]], ["Demo"])
        self.client.patch(f"{API}/portfolio/{linked.items.get().id}", {"name": "Edited in the new one"}, format="json")
        project.refresh_from_db()
        self.assertEqual(project.name, "Edited in the new one", "edit once, every portfolio changes")
        self.assertEqual(self.client.delete(f"{API}/portfolio/{project.id}?from={linked.id}").status_code, 204)
        self.assertEqual(self.shown_in(project), ["Alpha", "Beta"], "taking it out of the new portfolio leaves the others")

    def test_copying_without_sharing_is_unchanged_and_neither_flag_copies_no_projects(self):
        self.add(share_with=[self.b.id])
        both_off = self.client.post(f"{API}/tracks/", {"username": "owner", "name": "Blank-ish", "copy_from": self.a.id}, format="json")
        self.assertEqual(both_off.status_code, 201)
        self.assertEqual(both_off.json()["items"], [])
        copied = self.client.post(f"{API}/tracks/", {"username": "owner", "name": "Copied", "copy_from": self.a.id, "copy_portfolio_items": True}, format="json")
        self.assertEqual(copied.status_code, 201)
        self.assertEqual(Portfolio.objects.count(), 2, "a separate copy was made")
        self.assertEqual(Portfolio.objects.get(pk=copied.json()["items"][0]["id"]).tracks.count(), 1, "and the copy is shown in the new portfolio only")

    def test_asking_for_both_shares_rather_than_doubling_up(self):
        self.add()
        res = self.client.post(f"{API}/tracks/", {"username": "owner", "name": "Both", "copy_from": self.a.id, "copy_portfolio_items": True, "share_portfolio_items": True}, format="json")
        self.assertEqual(res.status_code, 201)
        self.assertEqual(Portfolio.objects.count(), 1)

    # --- the resume -------------------------------------------------------------------------------------

    @override_settings(SITE_URL="https://site.example")
    def test_a_shared_project_is_on_every_portfolios_resume_and_follows_the_ticks(self):
        project = self.add(share_with=[self.b.id], description="Shared words.")
        for track in (self.a, self.b):
            data = resume_builder.collect(track, "https://site.example")
            self.assertEqual([p["name"] for p in data["projects"]], ["Demo"], track.name)
            self.assertEqual(data["projects"][0]["description"], "Shared words.")
        self.assertEqual(resume_builder.collect(self.c, "https://site.example")["projects"], [])
        self.client.patch(f"{API}/portfolio/{project.id}", {"tracks": [self.a.id]}, format="json")
        self.assertEqual(resume_builder.collect(self.b, "https://site.example")["projects"], [])
        self.client.patch(f"{API}/portfolio/{project.id}", {"description": "Edited."}, format="json")
        self.assertEqual(resume_builder.collect(self.a, "https://site.example")["projects"][0]["description"], "Edited.")

    @override_settings(SITE_URL="https://site.example")
    def test_the_five_most_recent_rule_counts_shared_projects_like_any_other(self):
        for i in range(6):
            self.add(name=f"P{i}", share_with=[self.b.id] if i % 2 == 0 else None)
        self.assertEqual([p["name"] for p in resume_builder.collect(self.a, "https://site.example")["projects"]], ["P5", "P4", "P3", "P2", "P1"])
        self.assertEqual([p["name"] for p in resume_builder.collect(self.b, "https://site.example")["projects"]], ["P4", "P2", "P0"])


class ProjectSharingMigrationTests(TransactionTestCase):
    """Projects that exist before sharing keep showing in their portfolio, and the change can be undone."""

    def test_existing_projects_move_to_the_new_link_and_back(self):
        from django.db import connection
        from django.db.migrations.executor import MigrationExecutor
        before, after = [("portfolio", "0019_project_description")], [("portfolio", "0022_project_sharing_drop_old_link")]
        executor = MigrationExecutor(connection)
        executor.migrate(before)
        old = MigrationExecutor(connection).loader.project_state(before).apps
        user = old.get_model("auth", "User").objects.create(username="legacy")
        profile = old.get_model("portfolio", "Profile").objects.create(user=user, display_pic="https://x.example/p.png")
        Track, Project = old.get_model("portfolio", "PortfolioTrack"), old.get_model("portfolio", "Portfolio")
        one = Track.objects.create(profile=profile, name="One", slug="one")
        two = Track.objects.create(profile=profile, name="Two", slug="two")
        Project.objects.create(track=one, category="WebDev", name="In one", thumbnail="https://x.example/a.png", project_link="https://x.example/a")
        Project.objects.create(track=one, category="WebDev", name="Also in one", thumbnail="https://x.example/b.png", project_link="https://x.example/b")
        Project.objects.create(track=two, category="CAD", name="In two", thumbnail="https://x.example/c.png", project_link="https://x.example/c", description="Kept.")

        executor = MigrationExecutor(connection)
        executor.migrate(after)
        new = executor.loader.project_state(after).apps
        shown = {p.name: sorted(t.name for t in p.tracks.all()) for p in new.get_model("portfolio", "Portfolio").objects.all()}
        self.assertEqual(shown, {"In one": ["One"], "Also in one": ["One"], "In two": ["Two"]})
        self.assertEqual(new.get_model("portfolio", "Portfolio").objects.get(name="In two").description, "Kept.")

        executor = MigrationExecutor(connection)
        executor.migrate(before)
        back = executor.loader.project_state(before).apps
        homes = {p.name: p.track.name for p in back.get_model("portfolio", "Portfolio").objects.select_related("track")}
        self.assertEqual(homes, {"In one": "One", "Also in one": "One", "In two": "Two"})

        MigrationExecutor(connection).migrate(MigrationExecutor(connection).loader.graph.leaf_nodes())   # leave the schema as found


# ---------------------------------------------------------------- resume ------

import docx
import pymupdf

from portfolio import resume_builder, resume_mark, resume_photo

TODAY = date(2026, 10, 18)

ATS_DESIGNS = ["ats", "classic", "compact", "executive", "technical"]
GRAPHIC_DESIGNS = ["modern", "midnight", "banner", "timeline", "elegant"]
STANDARD_FONTS = ("Helvetica", "Times", "Courier")      # the PDF base fonts every parser and printer has


def make_jpeg(size=(300, 400), colour=(120, 80, 60)):
    from PIL import Image
    out = io.BytesIO()
    Image.new("RGB", size, colour).save(out, "JPEG")
    return out.getvalue()


def pictures(page):
    """The page's pictures other than the faint logo watermark every resume page carries."""
    mark_size = resume_mark.faint_logo()[1:]
    return [info for info in page.get_image_info() if (info["width"], info["height"]) != mark_size]


def has_watermark(page):
    mark_size = resume_mark.faint_logo()[1:]
    return sum((info["width"], info["height"]) == mark_size for info in page.get_image_info()) == 1


def pdf_text(blob):
    with pymupdf.open(stream=blob, filetype="pdf") as doc:
        return "\n".join(page.get_text() for page in doc)


class ResumePhotoTests(SimpleTestCase):
    """The owner's picture is a web address they typed, fetched by our server:
    only public http(s) pictures are accepted, and any problem means initials."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        import http.server
        import threading
        cls.hits = []
        jpeg = make_jpeg((300, 500))
        from PIL import Image
        import os
        huge = io.BytesIO()   # a perfectly good picture, but far over the size limit
        Image.frombytes("RGB", (1800, 1800), os.urandom(1800 * 1800 * 3)).save(huge, "PNG", compress_level=0)
        huge = huge.getvalue()
        assert len(huge) > resume_photo.MAX_BYTES

        class Handler(http.server.BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass

            def do_GET(self):
                cls.hits.append(self.path)
                routes = {
                    "/ok.jpg": (200, "image/jpeg", jpeg),
                    "/page.html": (200, "text/html", b"<html>not a picture</html>"),
                    "/big.png": (200, "image/png", huge),
                    "/disguised": (200, "text/html", jpeg),     # picture bytes, but not served as an image
                    "/broken.jpg": (200, "image/jpeg", b"definitely not a jpeg"),
                }
                if self.path == "/to-ok":
                    self.send_response(302); self.send_header("Location", "/ok.jpg"); self.end_headers(); return
                if self.path == "/to-private":    # "localhost" is this same server, but the guard (patched below) calls it private
                    self.send_response(302); self.send_header("Location", f"http://localhost:{cls.server.server_address[1]}/ok.jpg"); self.end_headers(); return
                if self.path == "/loop":
                    self.send_response(302); self.send_header("Location", "/loop"); self.end_headers(); return
                status, kind, body = routes.get(self.path, (404, "text/plain", b"no"))
                self.send_response(status); self.send_header("Content-Type", kind); self.send_header("Content-Length", str(len(body))); self.end_headers()
                try:
                    self.wfile.write(body)
                except (BrokenPipeError, ConnectionResetError):   # the fetcher gave up on a too-big picture, as it should
                    pass

        cls.server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        cls.base = f"http://127.0.0.1:{cls.server.server_address[1]}"
        threading.Thread(target=cls.server.serve_forever, daemon=True).start()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        super().tearDownClass()

    def setUp(self):
        cache.clear()
        self.hits.clear()
        # the test server lives on 127.0.0.1, so for these tests only that address counts as public
        patcher = mock.patch.object(resume_photo, "_is_public", side_effect=lambda host: host == "127.0.0.1")
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_a_good_picture_comes_back_square_and_jpeg(self):
        from PIL import Image
        raw = resume_photo.fetch_photo(f"{self.base}/ok.jpg")
        image = Image.open(io.BytesIO(raw))
        self.assertEqual((image.format, image.size), ("JPEG", (resume_photo.SIDE, resume_photo.SIDE)))

    def test_it_is_fetched_once_then_remembered_even_when_it_failed(self):
        resume_photo.fetch_photo(f"{self.base}/ok.jpg"); resume_photo.fetch_photo(f"{self.base}/ok.jpg")
        resume_photo.fetch_photo(f"{self.base}/gone.jpg"); resume_photo.fetch_photo(f"{self.base}/gone.jpg")
        self.assertEqual(sorted(self.hits), ["/gone.jpg", "/ok.jpg"])

    def test_a_thumbnail_comes_back_as_a_four_by_three_jpeg(self):
        from PIL import Image
        raw = resume_photo.fetch_thumbnail(f"{self.base}/ok.jpg")
        image = Image.open(io.BytesIO(raw))
        self.assertEqual((image.format, image.size), ("JPEG", resume_photo.THUMB_SIZE))
        self.assertIsNone(resume_photo.fetch_thumbnail(f"{self.base}/page.html"))
        self.assertIsNone(resume_photo.fetch_thumbnail(""))

    def test_several_thumbnails_are_fetched_together_each_once(self):
        found = resume_photo.fetch_thumbnails([f"{self.base}/ok.jpg", f"{self.base}/ok.jpg", f"{self.base}/gone.jpg", "", None])
        self.assertEqual(sorted(found), [f"{self.base}/gone.jpg", f"{self.base}/ok.jpg"])
        self.assertTrue(found[f"{self.base}/ok.jpg"])
        self.assertIsNone(found[f"{self.base}/gone.jpg"])
        self.assertEqual(sorted(self.hits), ["/gone.jpg", "/ok.jpg"])
        self.assertEqual(resume_photo.fetch_thumbnails([]), {})

    def test_a_photo_and_a_thumbnail_of_the_same_address_are_kept_apart(self):
        photo, thumb = resume_photo.fetch_photo(f"{self.base}/ok.jpg"), resume_photo.fetch_thumbnail(f"{self.base}/ok.jpg")
        self.assertNotEqual(photo, thumb)

    def test_bad_pictures_give_nothing_and_never_raise(self):
        for path in ("/page.html", "/big.png", "/disguised", "/broken.jpg", "/gone.jpg", "/loop"):
            self.assertIsNone(resume_photo.fetch_photo(self.base + path), path)
        self.assertIsNone(resume_photo.fetch_photo(""))
        self.assertIsNone(resume_photo.fetch_photo(None))

    def test_a_redirect_to_a_picture_is_followed_but_one_to_a_private_address_is_not(self):
        self.assertTrue(resume_photo.fetch_photo(f"{self.base}/to-ok"))
        self.hits.clear()
        self.assertIsNone(resume_photo.fetch_photo(f"{self.base}/to-private"))
        self.assertEqual(self.hits, ["/to-private"], "the private address was never contacted")
        self.assertTrue(resume_photo.fetch_photo(f"{self.base}/ok.jpg"), "while the same picture is fine by its public name")


class ResumePhotoGuardTests(SimpleTestCase):
    """The real address check, with nothing patched."""

    def setUp(self):
        cache.clear()

    def test_private_and_local_addresses_are_never_fetched(self):
        with mock.patch("urllib.request.OpenerDirector.open") as opened:
            for url in ("http://127.0.0.1/a.jpg", "http://localhost/a.jpg", "http://[::1]/a.jpg", "http://10.0.0.5/a.jpg",
                        "http://192.168.1.10/a.jpg", "http://172.16.0.1/a.jpg", "http://169.254.169.254/latest/meta-data/",
                        "http://0.0.0.0/a.jpg", "http://2130706433/a.jpg", "file:///etc/passwd", "ftp://example.com/a.jpg",
                        "javascript:alert(1)", "http:///a.jpg", "data:image/png;base64,AAAA"):
                self.assertIsNone(resume_photo.fetch_photo(url), url)
                self.assertIsNone(resume_photo.fetch_thumbnail(url), url)
            opened.assert_not_called()

    def test_a_name_that_does_not_resolve_is_not_public(self):
        self.assertFalse(resume_photo._is_public("no-such-host.invalid"))
        self.assertTrue(resume_photo._is_public("8.8.8.8"))


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"], SITE_URL="https://site.example")
class ResumeExportTests(APITestCase):
    """One resume per portfolio: owner-only PDF and Word, and a public read-only
    viewer that only ever hands out page images."""

    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user("owner", password="Old-Passw0rd-xyz", first_name="Sulaimon", last_name="Lawal")
        cls.profile = Profile.objects.create(user=cls.user, display_pic="https://img.example.com/p.png")
        cls.fs = cls.make_track("Full Stack", "full-stack", skill="Full Stack Developer", email="fs@example.com",
                                phone="+2348000000001", location="Abuja, Nigeria", years=4)
        cls.cad = cls.make_track("CAD Engineer", "cad-engineer", skill="CAD Engineer", email="cad@example.com",
                                 phone="+2348000000002", location="Kaduna, Nigeria", years=2)

    @classmethod
    def make_track(cls, name, slug, *, skill, email, phone, location, years):
        track = PortfolioTrack.objects.create(
            profile=cls.profile, name=name, slug=slug, display_pic="https://img.example.com/p.png",
            email=email, phone_number=phone, location=location,
            linkedin="https://www.linkedin.com/in/owner", github=f"https://github.com/owner-{slug}",
        )
        about = About.objects.create(track=track, skill=skill, description=f"Summary text for {name}.\nSecond line.",
                                     experience_base_years=years)
        return track

    def setUp(self):
        cache.clear()
        self.photo = mock.patch("portfolio.resume_graphics.fetch_photo", side_effect=lambda url: make_jpeg() if url else None)
        self.fetch_photo = self.photo.start()
        self.addCleanup(self.photo.stop)
        self.thumbs = mock.patch("portfolio.resume_graphics.fetch_thumbnails",
                                 side_effect=lambda urls: {u: make_jpeg((480, 360), (zlib.crc32(u.encode()) % 200, 90, 160)) for u in urls if u})
        self.fetch_thumbnails = self.thumbs.start()
        self.addCleanup(self.thumbs.stop)
        self.client = APIClient(HTTP_HOST="localhost")
        self.client.force_authenticate(user=self.user)
        about = self.fs.about
        self.job_old = Resume.objects.create(about=about, post="Robotic Instructor", company_name="Al Izzah",
                                             start_year=date(2022, 9, 12), end_year=date(2022, 10, 14),
                                             description="Taught robotics basics.\n- Built line-following robots.")
        self.job_now = Resume.objects.create(about=about, post="Full Stack Developer", company_name="Pio Tech",
                                             start_year=date(2024, 8, 5), end_year=date.today() + timedelta(days=400),
                                             description="Built web apps in Python and React.", certificate_link="https://c.example/job.png")
        self.school = School.objects.create(about=about, school_name="Ahmadu Bello University", degree="B.Eng. Mechanical Engineering",
                                            start_year=date(2020, 1, 6), end_year=date(2024, 10, 5))
        for skill, rating in (("Rust", 87), ("Python", 78), ("C++", 52), ("C", 20)):
            Proficiency.objects.create(about=about, skill_name=skill, skill_range=rating)
        self.shop = make_project(self.fs, category="WebDev", name="Shop", thumbnail="https://a.com/t.png",
                                             project_link="https://shop.example.com", description="A small online shop with a cart and card payments.")
        Proficiency.objects.create(about=self.cad.about, skill_name="Solid Works", skill_range=70)
        Resume.objects.create(about=self.cad.about, post="Draftsman", company_name="Drawing Co",
                              start_year=date(2023, 1, 2), end_year=date(2023, 6, 30), description="Drew parts.")

    def url(self, slug, tail):
        return f"{API}/resume-export/{slug}/{tail}"

    def pdf(self, slug="full-stack"):
        res = self.client.get(self.url(slug, "download/pdf"))
        self.assertEqual(res.status_code, 200, res.content[:200])
        return res

    # --- who can download ------------------------------------------------------

    def test_only_the_owner_can_download(self):
        anon = APIClient(HTTP_HOST="localhost")
        for kind in ("pdf", "docx"):
            self.assertIn(anon.get(self.url("full-stack", f"download/{kind}")).status_code, (401, 403), kind)
        self.assertEqual(self.client.get(self.url("full-stack", "download/pdf")).status_code, 200)

    def test_unknown_portfolio_and_format(self):
        self.assertEqual(self.client.get(self.url("nope", "download/pdf")).status_code, 404)
        self.assertEqual(self.client.get(self.url("full-stack", "download/exe")).status_code, 404)
        self.assertEqual(APIClient(HTTP_HOST="localhost").get(self.url("nope", "pages")).status_code, 404)

    def test_headers_and_file_names(self):
        res = self.pdf()
        self.assertEqual(res["Content-Type"], "application/pdf")
        self.assertEqual(res["Content-Disposition"], 'attachment; filename="Sulaimon_Lawal_Full_Stack_Resume.pdf"')
        self.assertEqual(res["Cache-Control"], "no-store")
        self.assertTrue(res.content.startswith(b"%PDF"))
        word = self.client.get(self.url("cad-engineer", "download/docx"))
        self.assertEqual(word["Content-Disposition"], 'attachment; filename="Sulaimon_Lawal_CAD_Engineer_Resume.docx"')
        self.assertIn("wordprocessingml", word["Content-Type"])

    def test_file_name_is_reduced_to_safe_characters(self):
        self.fs.name = 'Full "Stack"/../x;\r\ny'
        self.fs.save()
        disposition = self.pdf()["Content-Disposition"]
        self.assertRegex(disposition, r'^attachment; filename="[A-Za-z0-9_]+\.pdf"$')

    def test_the_browser_may_read_the_file_name(self):
        from django.conf import settings
        self.assertIn("Content-Disposition", settings.CORS_EXPOSE_HEADERS)

    # --- the PDF is ATS-friendly -------------------------------------------------

    def test_pdf_has_the_right_content_in_reading_order(self):
        text = pdf_text(self.pdf().content)
        for needle in ("Sulaimon Lawal", "Full Stack Developer | 4+ years of experience", "fs@example.com",
                       "+2348000000001", "Abuja, Nigeria", "https://site.example/full-stack",
                       "https://github.com/owner-full-stack", "Summary text for Full Stack."):
            self.assertIn(needle, text, needle)
        order = [text.index(h) for h in ("PROFESSIONAL SUMMARY", "SKILLS", "PROFESSIONAL EXPERIENCE", "EDUCATION", "PROJECTS", "CERTIFICATES")]
        self.assertEqual(order, sorted(order), "standard section headings, in the usual order")
        self.assertLess(text.index("Sulaimon Lawal"), text.index("PROFESSIONAL SUMMARY"))

    def test_pdf_experience_dates_bullets_and_newest_first(self):
        text = pdf_text(self.pdf().content)
        self.assertLess(text.index("Full Stack Developer | Pio Tech"), text.index("Robotic Instructor | Al Izzah"))
        self.assertIn("09/2022 \u2013 10/2022", text)
        self.assertIn("08/2024 \u2013 Present", text, "an end date still ahead reads Present")
        self.assertIn("Built line-following robots.", text)
        self.assertIn("\u2022", text)

    def test_pdf_education_skills_projects_certificates(self):
        text = pdf_text(self.pdf().content)
        self.assertIn("B.Eng. Mechanical Engineering | Ahmadu Bello University", text)
        self.assertIn("2020 \u2013 2024", text)
        for line in ("Expert: Rust", "Advanced: Python", "Intermediate: C++", "Beginner: C"):
            self.assertIn(line, text, line)
        self.assertLess(text.index("Expert: Rust"), text.index("Advanced: Python"))
        self.assertIn("Shop (WebDev): https://shop.example.com", text)
        self.assertIn("Full Stack Developer, Pio Tech: https://c.example/job.png", text)

    def test_pdf_is_plain_enough_for_an_ats(self):
        with pymupdf.open(stream=self.pdf().content, filetype="pdf") as doc:
            self.assertEqual(doc.metadata["title"], "Sulaimon Lawal - Resume")
            self.assertEqual(doc.metadata["author"], "Sulaimon Lawal")
            fonts = set()
            for page in doc:
                self.assertEqual(pictures(page), [], "no pictures of its own: nothing for a parser to trip on")
                self.assertTrue(has_watermark(page), "just the faint logo behind the text")
                self.assertEqual(page.find_tables().tables, [], "no tables")
                fonts |= {f[3] for f in page.get_fonts()}
                lefts = [round(b[0]) for b in page.get_text("blocks") if b[4].strip()]
                self.assertLess(max(lefts) - min(lefts), 30, "a single column: every block starts at the left margin")
            self.assertTrue(fonts and all("Helvetica" in f for f in fonts), fonts)
            self.assertTrue(doc[0].get_text().strip(), "real, selectable text")
            self.assertTrue(any(link["kind"] == pymupdf.LINK_URI for link in doc[0].get_links()), "links stay clickable")

    def test_pdf_survives_awkward_text(self):
        self.job_now.description = "R&D <b>lead</b> \u2013 caf\u00e9 \u1ecc l\u00e1 \u201cquoted\u201d \u4e2d\u6587"
        self.job_now.save()
        self.user.last_name = "\u1ecc l\u00e1d\u00e9j\u00ed"
        self.user.save()
        text = pdf_text(self.pdf().content)
        self.assertIn("R&D <b>lead</b>", text, "markup in the text is shown literally, not interpreted")
        self.assertIn("caf\u00e9", text)
        self.assertIn("Sulaimon", text)

    def test_empty_portfolio_still_makes_a_valid_resume(self):
        blank = PortfolioTrack.objects.create(profile=self.profile, name="Blank", slug="blank", display_pic="https://img.example.com/p.png")
        About.objects.create(track=blank, skill="", description="")
        text = pdf_text(self.client.get(self.url("blank", "download/pdf")).content)
        self.assertIn("Sulaimon Lawal", text)
        for heading in ("PROFESSIONAL EXPERIENCE", "EDUCATION", "SKILLS", "PROJECTS", "CERTIFICATES", "PROFESSIONAL SUMMARY"):
            self.assertNotIn(heading, text, "empty sections are left out")
        self.assertEqual(self.client.get(self.url("blank", "download/docx")).status_code, 200)

    def test_a_long_resume_runs_onto_more_pages(self):
        for i in range(60):
            Resume.objects.create(about=self.fs.about, post=f"Role {i}", company_name=f"Company {i}", start_year=date(2010, 1, 1),
                                  end_year=date(2011, 1, 1), description="Did a thing.\nDid another thing.")
        with pymupdf.open(stream=self.pdf().content, filetype="pdf") as doc:
            self.assertGreater(len(doc), 2)
            self.assertIn("Role 59", "".join(p.get_text() for p in doc))

    # --- Word ---------------------------------------------------------------------

    def test_word_file_uses_real_headings_and_bullets(self):
        res = self.client.get(self.url("full-stack", "download/docx"))
        document = docx.Document(io.BytesIO(res.content))
        headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
        self.assertEqual(headings, ["PROFESSIONAL SUMMARY", "SKILLS", "PROFESSIONAL EXPERIENCE", "EDUCATION", "PROJECTS", "CERTIFICATES"])
        bullets = [p.text for p in document.paragraphs if p.style.name == "List Bullet"]
        self.assertIn("Built line-following robots.", bullets)
        text = "\n".join(p.text for p in document.paragraphs)
        for needle in ("Sulaimon Lawal", "fs@example.com", "+2348000000001", "Abuja, Nigeria", "08/2024 \u2013 Present",
                       "B.Eng. Mechanical Engineering | Ahmadu Bello University", "Expert: Rust"):
            self.assertIn(needle, text, needle)
        self.assertEqual(document.core_properties.title, "Sulaimon Lawal - Resume")
        self.assertEqual(len(document.tables), 0, "no tables")
        self.assertTrue(any("hyperlink" in p._p.xml for p in document.paragraphs), "links are real hyperlinks")

    # --- one resume per portfolio -----------------------------------------------------

    def test_each_portfolio_gets_its_own_resume(self):
        fs, cad = pdf_text(self.pdf("full-stack").content), pdf_text(self.pdf("cad-engineer").content)
        self.assertIn("fs@example.com", fs); self.assertNotIn("cad@example.com", fs)
        self.assertIn("cad@example.com", cad); self.assertNotIn("fs@example.com", cad)
        self.assertIn("Kaduna, Nigeria", cad); self.assertNotIn("Abuja", cad)
        self.assertIn("CAD Engineer | 2+ years of experience", cad)
        self.assertIn("Draftsman | Drawing Co", cad); self.assertNotIn("Draftsman", fs)
        self.assertIn("Solid Works", cad); self.assertNotIn("Rust", cad)
        self.assertIn("https://site.example/cad-engineer", cad)

    def test_editing_a_portfolio_changes_its_resume_and_not_the_other(self):
        before = self.pdf("full-stack").content
        other = self.pdf("cad-engineer").content
        self.client.patch(f"{API}/tracks/{self.fs.id}", {"email": "new@example.com", "location": "Lagos, Nigeria"}, format="json")
        after = pdf_text(self.pdf("full-stack").content)
        self.assertIn("new@example.com", after); self.assertIn("Lagos, Nigeria", after)
        self.assertNotEqual(before, self.pdf("full-stack").content)
        self.assertEqual(pdf_text(other), pdf_text(self.pdf("cad-engineer").content))

    def test_results_are_cached_until_the_content_changes(self):
        with mock.patch.object(resume_builder, "render_pdf", wraps=resume_builder.render_pdf) as render:
            self.pdf(); self.pdf(); self.pdf()
            self.assertEqual(render.call_count, 1)
            Proficiency.objects.create(about=self.fs.about, skill_name="Go", skill_range=65)
            self.assertIn("Go", pdf_text(self.pdf().content))
            self.assertEqual(render.call_count, 2)

    # --- the public read-only viewer --------------------------------------------------

    def test_viewer_is_public_and_serves_page_images_only(self):
        anon = APIClient(HTTP_HOST="localhost")
        info = anon.get(self.url("full-stack", "pages")).json()
        self.assertEqual((info["name"], info["portfolio"]), ("Sulaimon Lawal", "Full Stack"))
        self.assertGreaterEqual(info["pages"], 1)
        image = anon.get(self.url("full-stack", "pages/1.png"))
        self.assertEqual(image.status_code, 200)
        self.assertEqual(image["Content-Type"], "image/png")
        self.assertTrue(image.content.startswith(b"\x89PNG"))
        self.assertEqual(image["Cache-Control"], "no-store")
        self.assertEqual(anon.get(self.url("full-stack", f"pages/{info['pages'] + 1}.png")).status_code, 404)
        self.assertEqual(anon.get(self.url("full-stack", "pages/0.png")).status_code, 404)

    def test_viewer_pages_are_pictures_with_no_text_to_copy(self):
        blob = APIClient(HTTP_HOST="localhost").get(self.url("full-stack", "pages/1.png")).content
        self.assertEqual(pymupdf.open(stream=blob, filetype="png").page_count, 1)
        with pymupdf.open(stream=blob, filetype="png") as picture:
            self.assertEqual(picture[0].get_text().strip(), "", "an image has no text layer")

    def test_viewer_differs_per_portfolio(self):
        anon = APIClient(HTTP_HOST="localhost")
        self.assertNotEqual(anon.get(self.url("full-stack", "pages/1.png")).content, anon.get(self.url("cad-engineer", "pages/1.png")).content)

    def test_viewer_is_rate_limited(self):
        with mock.patch.dict(ScopedRateThrottle.THROTTLE_RATES, {"resume_view": "3/hour"}):
            anon = APIClient(HTTP_HOST="localhost")
            codes = [anon.get(self.url("full-stack", "pages")).status_code for _ in range(5)]
        self.assertEqual(codes, [200, 200, 200, 429, 429])

    def test_viewer_cannot_be_used_to_get_the_files(self):
        anon = APIClient(HTTP_HOST="localhost")
        self.assertIn(anon.get(self.url("full-stack", "download/pdf")).status_code, (401, 403))
        self.assertEqual(anon.get(self.url("full-stack", "pages/1.pdf")).status_code, 404)

    # --- the designs -----------------------------------------------------------------------

    def get(self, slug, tail, template=None, client=None):
        suffix = f"?template={template}" if template else ""
        return (client or self.client).get(self.url(slug, tail) + suffix)

    def test_the_design_list_is_public_and_has_five_of_each_kind(self):
        res = APIClient(HTTP_HOST="localhost").get(f"{API}/resume-export/templates")
        self.assertEqual(res.status_code, 200)
        designs = res.json()
        self.assertEqual([d["key"] for d in designs if d["group"] == "ats"], ATS_DESIGNS)
        self.assertEqual([d["key"] for d in designs if d["group"] == "graphic"], GRAPHIC_DESIGNS)
        self.assertEqual([d["key"] for d in designs if d["default"]], ["ats"])
        for design in designs:
            self.assertTrue(design["label"] and design["blurb"], design)
            self.assertEqual(design["formats"], ["pdf", "docx"] if design["group"] == "ats" else ["pdf"], design["key"])

    def test_every_ats_design_is_plain_text_a_parser_can_read(self):
        for key in ATS_DESIGNS:
            with self.subTest(key):
                res = self.get("full-stack", "download/pdf", key)
                self.assertEqual(res.status_code, 200, res.content[:200])
                with pymupdf.open(stream=res.content, filetype="pdf") as doc:
                    text = "".join(page.get_text() for page in doc)
                    for needle in ("Sulaimon Lawal", "fs@example.com", "+2348000000001", "Abuja, Nigeria", "Rust", "Pio Tech",
                                   "Ahmadu Bello University", "Shop", "https://site.example/full-stack"):
                        self.assertIn(needle, text, f"{key}: {needle}")
                    order = [text.index(h) for h in ("PROFESSIONAL SUMMARY", "SKILLS", "PROFESSIONAL EXPERIENCE", "EDUCATION", "PROJECTS")]
                    self.assertEqual(order, sorted(order), key)
                    fonts = set()
                    for page in doc:
                        self.assertEqual(pictures(page), [], f"{key}: no pictures of its own")
                        self.assertTrue(has_watermark(page), f"{key}: the faint logo is behind the text")
                        self.assertEqual(page.find_tables().tables, [], f"{key}: no tables")
                        fonts |= {f[3] for f in page.get_fonts()}
                        boxes = [b for b in page.get_text("blocks") if b[4].strip()]
                        for i, a in enumerate(boxes):
                            for b in boxes[i + 1:]:
                                stacked = min(a[3], b[3]) - max(a[1], b[1]) > 3
                                disjoint = a[2] < b[0] or b[2] < a[0]
                                self.assertFalse(stacked and disjoint, f"{key}: two columns side by side: {a[4]!r} / {b[4]!r}")
                    self.assertTrue(fonts and all(f.startswith(STANDARD_FONTS) for f in fonts), f"{key}: {fonts}")
                    self.assertTrue(any(l["kind"] == pymupdf.LINK_URI for l in doc[0].get_links()), f"{key}: links clickable")
                    self.assertEqual(doc.metadata["title"], "Sulaimon Lawal - Resume")

    def test_every_ats_design_has_a_word_version_with_real_structure(self):
        for key in ATS_DESIGNS:
            with self.subTest(key):
                res = self.get("full-stack", "download/docx", key)
                self.assertEqual(res.status_code, 200)
                document = docx.Document(io.BytesIO(res.content))
                headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
                self.assertEqual(headings, ["PROFESSIONAL SUMMARY", "SKILLS", "PROFESSIONAL EXPERIENCE", "EDUCATION", "PROJECTS", "CERTIFICATES"], key)
                self.assertIn("Built line-following robots.", [p.text for p in document.paragraphs if p.style.name == "List Bullet"])
                self.assertEqual((len(document.tables), len(document.inline_shapes)), (0, 0), f"{key}: no tables or pictures")
                self.assertTrue(any("hyperlink" in p._p.xml for p in document.paragraphs), key)

    def test_the_ats_designs_look_different_from_each_other(self):
        pdfs = {key: self.get("full-stack", "download/pdf", key).content for key in ATS_DESIGNS}
        self.assertEqual(len(set(pdfs.values())), 5)
        words = {key: self.get("full-stack", "download/docx", key).content for key in ATS_DESIGNS}
        self.assertEqual(len(set(words.values())), 5)
        with pymupdf.open(stream=pdfs["classic"], filetype="pdf") as classic, pymupdf.open(stream=pdfs["technical"], filetype="pdf") as technical:
            self.assertTrue(any("Times" in f[3] for f in classic[0].get_fonts()))
            self.assertTrue(any("Courier" in f[3] for f in technical[0].get_fonts()))

    def test_every_graphic_design_is_a_pdf_with_the_photo(self):
        for key in GRAPHIC_DESIGNS:
            with self.subTest(key):
                res = self.get("full-stack", "download/pdf", key)
                self.assertEqual(res.status_code, 200, res.content[:200])
                with pymupdf.open(stream=res.content, filetype="pdf") as doc:
                    self.assertEqual(len(pictures(doc[0])), 2, f"{key}: the photo and the project's thumbnail are on page one")
                    text = doc[0].get_text()
                    for needle in ("Sulaimon Lawal", "fs@example.com", "Rust", "Pio Tech", "Ahmadu Bello University", "Shop",
                                   "A small online shop with a cart and card payments.", "shop.example.com"):
                        self.assertIn(needle, text, f"{key}: {needle}")
                    self.assertTrue(any(l["kind"] == pymupdf.LINK_URI for l in doc[0].get_links()), f"{key}: links clickable")

    def test_graphic_designs_put_the_photo_in_different_places(self):
        spots = {}
        for key in GRAPHIC_DESIGNS:
            with pymupdf.open(stream=self.get("full-stack", "download/pdf", key).content, filetype="pdf") as doc:
                boxes = [pymupdf.Rect(info["bbox"]) for info in pictures(doc[0])]
                box = next(b for b in boxes if abs(b.width / b.height - 4 / 3) > 0.05)   # the photo, not a 4:3 thumbnail
                spots[key] = (round(box.x0 / doc[0].rect.width, 1), round(box.y0 / doc[0].rect.height, 1), round(box.width))
        self.assertEqual(len(set(spots.values())), 5, spots)

    def test_graphic_designs_are_not_available_as_word(self):
        for key in GRAPHIC_DESIGNS:
            res = self.get("full-stack", "download/docx", key)
            self.assertEqual(res.status_code, 400, key)
            self.assertIn("template", res.json())

    def test_no_photo_means_initials_in_the_same_spot(self):
        self.fetch_photo.side_effect = lambda url: None
        for key in GRAPHIC_DESIGNS:
            with self.subTest(key):
                res = self.get("full-stack", "download/pdf", key)
                self.assertEqual(res.status_code, 200)
                with pymupdf.open(stream=res.content, filetype="pdf") as doc:
                    self.assertEqual(len(pictures(doc[0])), 1, f"{key}: only the project thumbnail is left")
                    self.assertIn("SL", doc[0].get_text().split(), key)

    def test_unknown_design_is_rejected_everywhere(self):
        for tail in ("download/pdf", "download/docx", "pages", "pages/1.png"):
            res = self.get("full-stack", tail, "fancy")
            self.assertEqual(res.status_code, 400, tail)
        self.assertEqual(self.get("full-stack", "download/pdf", "../etc").status_code, 400)

    def test_only_the_owner_can_download_any_design(self):
        anon = APIClient(HTTP_HOST="localhost")
        for key in ATS_DESIGNS + GRAPHIC_DESIGNS:
            self.assertIn(self.get("full-stack", "download/pdf", key, client=anon).status_code, (401, 403), key)

    def test_file_names_carry_the_design(self):
        self.assertEqual(self.get("full-stack", "download/pdf")["Content-Disposition"], 'attachment; filename="Sulaimon_Lawal_Full_Stack_Resume.pdf"')
        self.assertEqual(self.get("full-stack", "download/pdf", "ats")["Content-Disposition"], 'attachment; filename="Sulaimon_Lawal_Full_Stack_Resume.pdf"')
        self.assertEqual(self.get("full-stack", "download/pdf", "elegant")["Content-Disposition"], 'attachment; filename="Sulaimon_Lawal_Full_Stack_Resume_Elegant.pdf"')
        self.assertEqual(self.get("cad-engineer", "download/docx", "classic")["Content-Disposition"], 'attachment; filename="Sulaimon_Lawal_CAD_Engineer_Resume_Classic.docx"')

    def test_every_design_survives_awkward_text_an_empty_portfolio_and_a_long_history(self):
        self.job_now.description = "R&D <b>lead</b> \u2013 caf\u00e9 \u1ecc l\u00e1 \u201cquoted\u201d \u4e2d\u6587"
        self.job_now.save()
        self.user.first_name, self.user.last_name = "Chukwuemeka-Oluwaseun", "\u1ecc l\u00e1d\u00e9j\u00ed-Adeyemi-Okonkwo"
        self.user.save()
        blank = PortfolioTrack.objects.create(profile=self.profile, name="Blank", slug="blank", display_pic="")
        About.objects.create(track=blank, skill="", description="")
        for i in range(40):
            Resume.objects.create(about=self.fs.about, post=f"Role {i}", company_name=f"Company {i}", start_year=date(2010, 1, 1),
                                  end_year=date(2011, 1, 1), description="Did a thing.\nDid another thing.")
        for key in ATS_DESIGNS + GRAPHIC_DESIGNS:
            with self.subTest(key):
                text = pdf_text(self.get("full-stack", "download/pdf", key).content)
                self.assertIn("R&D <b>lead</b>", text)
                self.assertIn("Role 39", text, "a long history is not cut off")
                self.assertEqual(self.get("blank", "download/pdf", key).status_code, 200)

    def test_each_design_is_cached_separately_until_the_content_changes(self):
        with mock.patch.object(resume_builder, "render_pdf", wraps=resume_builder.render_pdf) as render:
            for key in ("ats", "elegant", "ats", "elegant", "elegant"):
                self.get("full-stack", "download/pdf", key)
            self.assertEqual(render.call_count, 2, "one render per design")
            self.client.patch(f"{API}/tracks/{self.fs.id}", {"display_pic": "https://img.example.com/other.png"}, format="json")
            self.get("full-stack", "download/pdf", "elegant")
            self.assertEqual(render.call_count, 3, "a new photo address is new content")

    # --- project cards: thumbnail, short description, link -----------------------------------

    def card_pdf(self, key="modern"):
        res = self.get("full-stack", "download/pdf", key)
        self.assertEqual(res.status_code, 200, res.content[:200])
        return res.content

    def test_graphic_designs_show_each_project_with_thumbnail_description_and_link(self):
        make_project(self.fs, category="Robotics", name="Fence Lamp", thumbnail="https://a.com/lamp.png",
                                 project_link="https://lamp.example.com", description="Lights a fence line at dusk.")
        for key in GRAPHIC_DESIGNS:
            with self.subTest(key):
                with pymupdf.open(stream=self.card_pdf(key), filetype="pdf") as doc:
                    text = "".join(page.get_text() for page in doc)
                    images = sum(len(pictures(page)) for page in doc)
                    uris = {l["uri"] for page in doc for l in page.get_links() if l["kind"] == pymupdf.LINK_URI}
                    self.assertEqual(images, 3, f"{key}: the photo and two thumbnails")
                    for needle in ("Shop", "A small online shop with a cart and card payments.", "Fence Lamp", "Lights a fence line at dusk."):
                        self.assertIn(needle, text, f"{key}: {needle}")
                    self.assertIn("https://shop.example.com", uris, f"{key}: the link is clickable")
                    self.assertIn("https://lamp.example.com", uris, key)

    def test_the_description_is_optional_on_the_card(self):
        self.shop.description = ""
        self.shop.save()
        for key in GRAPHIC_DESIGNS:
            with self.subTest(key):
                with pymupdf.open(stream=self.card_pdf(key), filetype="pdf") as doc:
                    text = doc[0].get_text()
                    self.assertIn("Shop", text)
                    self.assertIn("shop.example.com", text)
                    self.assertNotIn("None", text)
                    self.assertEqual(len(pictures(doc[0])), 2)

    def test_a_project_whose_thumbnail_cannot_be_fetched_gets_a_plain_tile(self):
        self.fetch_thumbnails.side_effect = lambda urls: {u: None for u in urls}
        for key in GRAPHIC_DESIGNS:
            with self.subTest(key):
                with pymupdf.open(stream=self.card_pdf(key), filetype="pdf") as doc:
                    text = doc[0].get_text()
                    self.assertEqual(len(pictures(doc[0])), 1, f"{key}: just the photo")
                    self.assertIn("A small online shop", text)
                    self.assertIn("S", text.split(), f"{key}: the tile carries the project's first letter")

    def test_a_resume_carries_the_five_most_recent_projects_newest_first(self):
        for i in range(8):
            make_project(self.fs, category="WebDev", name=f"Project {i}", thumbnail=f"https://a.com/{i}.png",
                                     project_link=f"https://p{i}.example.com", description="Does a thing. " * 10)
        # 9 projects now (Shop is the oldest); the five added last are Project 7 .. Project 3
        for key in ATS_DESIGNS + GRAPHIC_DESIGNS:
            with self.subTest(key):
                with pymupdf.open(stream=self.get("full-stack", "download/pdf", key).content, filetype="pdf") as doc:
                    text = "".join(page.get_text() for page in doc)
                for i in range(3, 8):
                    self.assertIn(f"Project {i}", text, f"{key}: project {i} is one of the five newest")
                for old in ("Project 2", "Project 1", "Project 0", "Shop"):
                    self.assertNotIn(old, text, f"{key}: {old} is older than the five newest")
                self.assertLess(text.index("Project 7"), text.index("Project 3"), f"{key}: newest first")
        word = docx.Document(io.BytesIO(self.get("full-stack", "download/docx").content))
        joined = "\n".join(p.text for p in word.paragraphs)
        self.assertIn("Project 7", joined); self.assertNotIn("Project 2", joined); self.assertNotIn("Shop", joined)
        # the website still lists every one of them
        tracks = APIClient(HTTP_HOST="localhost").get(f"{API}/tracks/?username=owner").json()
        self.assertEqual(len(next(t for t in tracks if t["slug"] == "full-stack")["items"]), 9)

    def test_five_or_fewer_projects_all_appear(self):
        for i in range(4):
            make_project(self.fs, category="WebDev", name=f"Project {i}", thumbnail=f"https://a.com/{i}.png",
                                     project_link=f"https://p{i}.example.com", description="Does a thing. " * 10)
        for key in ATS_DESIGNS + GRAPHIC_DESIGNS:
            with self.subTest(key):
                text = pdf_text(self.get("full-stack", "download/pdf", key).content)
                for name in ("Shop", "Project 0", "Project 1", "Project 2", "Project 3"):
                    self.assertIn(name, text, f"{key}: {name}")
        self.assertEqual(len(resume_builder.collect(self.cad, "https://site.example")["projects"]), 0)

    def test_the_cap_follows_new_projects_and_cached_files_do_not_hide_it(self):
        before = pdf_text(self.get("full-stack", "download/pdf").content)
        self.assertIn("Shop", before)
        for i in range(5):
            make_project(self.fs, category="WebDev", name=f"Fresh {i}", thumbnail=f"https://a.com/f{i}.png",
                                     project_link=f"https://f{i}.example.com")
        after = pdf_text(self.get("full-stack", "download/pdf").content)
        self.assertNotIn("Shop", after, "five newer projects push the oldest off")
        self.assertIn("Fresh 4", after)

    def test_ats_designs_show_the_description_and_link_but_no_thumbnail(self):
        for key in ATS_DESIGNS:
            with self.subTest(key):
                with pymupdf.open(stream=self.get("full-stack", "download/pdf", key).content, filetype="pdf") as doc:
                    text = "".join(page.get_text() for page in doc)
                    self.assertIn("Shop (WebDev): https://shop.example.com", text, key)
                    self.assertIn("A small online shop with a cart and card payments.", text, key)
                    self.assertTrue(all(pictures(page) == [] for page in doc), f"{key}: no thumbnail")
                word = docx.Document(io.BytesIO(self.get("full-stack", "download/docx", key).content))
                joined = "\n".join(p.text for p in word.paragraphs)
                self.assertIn("A small online shop with a cart and card payments.", joined, key)
                self.assertEqual(len(word.inline_shapes), 0, key)
        self.fetch_thumbnails.assert_not_called()

    def test_ats_designs_still_read_fine_without_a_description(self):
        self.shop.description = ""
        self.shop.save()
        text = pdf_text(self.pdf().content)
        self.assertIn("Shop (WebDev): https://shop.example.com", text)
        self.assertNotIn("None", text)

    def test_changing_a_description_changes_the_resume(self):
        before = self.card_pdf("elegant")
        self.client.patch(f"{API}/portfolio/{self.shop.id}", {"description": "Now with gift cards."}, format="json")
        self.assertIn("Now with gift cards.", pdf_text(self.card_pdf("elegant")))
        self.assertNotEqual(before, self.card_pdf("elegant"))
        self.assertIn("Now with gift cards.", pdf_text(self.pdf().content))

    def test_project_description_is_optional_short_and_copied(self):
        body = {"track": self.fs.id, "category": "WebDev", "name": "Blog", "thumbnail": "https://a.com/b.png", "project_link": "https://blog.example.com"}
        res = self.client.post(f"{API}/portfolio/", body, format="json")
        self.assertEqual((res.status_code, res.json()["description"]), (201, ""), "left out is fine")
        res = self.client.post(f"{API}/portfolio/", {**body, "name": "Blog 2", "description": ""}, format="json")
        self.assertEqual(res.status_code, 201)
        res = self.client.post(f"{API}/portfolio/", {**body, "name": "Blog 3", "description": "x" * 160}, format="json")
        self.assertEqual(res.status_code, 201, "160 characters is the limit")
        res = self.client.post(f"{API}/portfolio/", {**body, "name": "Blog 4", "description": "x" * 161}, format="json")
        self.assertEqual(res.status_code, 400)
        self.assertIn("description", res.json())
        self.assertEqual(self.client.patch(f"{API}/portfolio/{self.shop.id}", {"description": ""}, format="json").status_code, 200)
        self.shop.refresh_from_db()
        self.assertEqual(self.shop.description, "")
        # in the public payload
        self.client.patch(f"{API}/portfolio/{self.shop.id}", {"description": "Kept."}, format="json")
        tracks = APIClient(HTTP_HOST="localhost").get(f"{API}/tracks/?username=owner").json()
        items = next(t for t in tracks if t["slug"] == "full-stack")["items"]
        self.assertEqual({i["name"]: i["description"] for i in items}["Shop"], "Kept.")
        # duplicating a portfolio carries it over, and the copy then stands alone
        res = self.client.post(f"{API}/tracks/", {"username": "owner", "name": "Copy", "copy_from": self.fs.id, "copy_portfolio_items": True}, format="json")
        self.assertEqual(res.status_code, 201, res.content)
        copy = PortfolioTrack.objects.get(pk=res.json()["id"])
        self.assertEqual(copy.items.get(name="Shop").description, "Kept.")
        self.client.patch(f"{API}/portfolio/{copy.items.get(name='Shop').id}", {"description": "Changed in the copy."}, format="json")
        self.shop.refresh_from_db()
        self.assertEqual(self.shop.description, "Kept.")

    # --- the logo watermark ------------------------------------------------------------------------

    def test_every_page_of_every_design_carries_the_faint_logo(self):
        for i in range(30):
            make_project(self.fs, category="WebDev", name=f"Filler {i}", thumbnail=f"https://a.com/f{i}.png", project_link=f"https://f{i}.example.com")
        for i in range(25):
            Resume.objects.create(about=self.fs.about, post=f"Role {i}", company_name=f"Company {i}", start_year=date(2010, 1, 1),
                                  end_year=date(2011, 1, 1), description="Did a thing.\nDid another thing.")
        for key in ATS_DESIGNS + GRAPHIC_DESIGNS:
            with self.subTest(key):
                with pymupdf.open(stream=self.get("full-stack", "download/pdf", key).content, filetype="pdf") as doc:
                    self.assertGreaterEqual(len(doc), 2)
                    for number, page in enumerate(doc, 1):
                        self.assertTrue(has_watermark(page), f"{key}: page {number}")

    def test_the_watermark_is_a_picture_behind_the_text_and_changes_nothing_in_the_text(self):
        for key in ATS_DESIGNS + GRAPHIC_DESIGNS:
            with self.subTest(key):
                with_mark = self.get("full-stack", "download/pdf", key).content
                cache.clear()
                with mock.patch.object(resume_mark, "draw_on_page"):
                    without = self.get("full-stack", "download/pdf", key).content
                cache.clear()
                self.assertNotEqual(with_mark, without)
                self.assertEqual(pdf_text(with_mark), pdf_text(without), f"{key}: the text layer and its order are untouched")
                with pymupdf.open(stream=with_mark, filetype="pdf") as doc:
                    info = next(i for i in doc[0].get_image_info() if (i["width"], i["height"]) == resume_mark.faint_logo()[1:])
                    box = pymupdf.Rect(info["bbox"])
                    self.assertAlmostEqual((box.x0 + box.x1) / 2, doc[0].rect.width / 2, delta=2)
                    self.assertAlmostEqual((box.y0 + box.y1) / 2, doc[0].rect.height / 2, delta=2)
                    self.assertGreater(box.width, doc[0].rect.width * 0.5, "big")
                    self.assertNotEqual(tuple(round(v, 3) for v in info["transform"][1:3]), (0.0, 0.0), "slanted")

    def test_the_logo_is_faint_and_grey_so_it_cannot_drown_the_text(self):
        data, width, height = resume_mark.faint_logo()
        from PIL import Image
        image = Image.open(io.BytesIO(data))
        self.assertEqual((image.mode, image.size), ("RGBA", (width, height)))
        alphas = list(image.getchannel("A").getdata())
        self.assertLessEqual(max(alphas), 30, "at most about 12% opaque")
        self.assertGreater(max(alphas), 0)
        self.assertEqual({px[:3] for px in image.getdata() if px[3]}, {(110, 114, 123)}, "one mid-grey")

    def test_word_files_have_the_logo_in_the_header_behind_the_text(self):
        for key in ATS_DESIGNS:
            with self.subTest(key):
                document = docx.Document(io.BytesIO(self.get("full-stack", "download/docx", key).content))
                header = document.sections[0].header
                self.assertFalse(header.is_linked_to_previous)
                xml = header._element.xml
                self.assertIn("wp:anchor", xml)
                self.assertIn('behindDoc="1"', xml, "behind the text")
                self.assertIn("<wp:wrapNone/>", xml, "text does not flow around it")
                self.assertIn('rot="-840000"', xml, "slanted 14 degrees")
                self.assertIn("a:blip", xml)
                self.assertEqual(len(document.inline_shapes), 0, "nothing inline in the resume body")
                self.assertEqual(len(document.tables), 0)
                headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
                self.assertEqual(headings[0], "PROFESSIONAL SUMMARY", "the body and its headings are as before")
                images = [part for part in header.part.related_parts.values() if getattr(part, "content_type", "").startswith("image/")]
                self.assertEqual(len(images), 1)

    def test_previews_and_the_visitor_viewer_show_the_watermark_too(self):
        anon = APIClient(HTTP_HOST="localhost")
        for key in ("ats", "modern"):
            with self.subTest(key):
                cache.clear()
                with_mark = anon.get(self.url("full-stack", "pages/1.png") + f"?template={key}").content
                thumb = anon.get(self.url("full-stack", "pages/1.png") + f"?template={key}&thumb=1").content
                cache.clear()
                with mock.patch.object(resume_mark, "draw_on_page"):
                    without = anon.get(self.url("full-stack", "pages/1.png") + f"?template={key}").content
                    thumb_without = anon.get(self.url("full-stack", "pages/1.png") + f"?template={key}&thumb=1").content
                cache.clear()
                self.assertNotEqual(with_mark, without, "the page image the visitor and the owner's preview show carries the mark")
                self.assertNotEqual(thumb, thumb_without, "and so do the thumbnails")

    # --- picking a design in the viewer -----------------------------------------------------

    def test_the_viewer_defaults_to_a_designed_layout_and_can_show_any_design(self):
        anon = APIClient(HTTP_HOST="localhost")
        self.assertEqual(anon.get(self.url("full-stack", "pages")).json()["template"], "modern")
        pages = {key: self.get("full-stack", "pages/1.png", key, client=anon).content for key in ATS_DESIGNS + GRAPHIC_DESIGNS}
        self.assertEqual(len(set(pages.values())), 10)
        for blob in pages.values():
            self.assertTrue(blob.startswith(b"\x89PNG"))

    def test_design_thumbnails_are_small_first_pages(self):
        anon = APIClient(HTTP_HOST="localhost")
        full = anon.get(self.url("full-stack", "pages/1.png") + "?template=elegant").content
        for key in ATS_DESIGNS + GRAPHIC_DESIGNS:
            res = anon.get(self.url("full-stack", "pages/1.png") + f"?template={key}&thumb=1")
            self.assertEqual((res.status_code, res["Content-Type"]), (200, "image/png"), key)
            self.assertTrue(res.content.startswith(b"\x89PNG"))
        thumb = anon.get(self.url("full-stack", "pages/1.png") + "?template=elegant&thumb=1").content
        self.assertLess(len(thumb), len(full))
        self.assertEqual(anon.get(self.url("full-stack", "pages/2.png") + "?template=elegant&thumb=1").status_code, 404)
        self.assertEqual(anon.get(self.url("nope", "pages/1.png") + "?thumb=1").status_code, 404)

    def test_the_viewer_still_hands_out_pictures_only_for_every_design(self):
        anon = APIClient(HTTP_HOST="localhost")
        for key in ATS_DESIGNS + GRAPHIC_DESIGNS:
            blob = self.get("full-stack", "pages/1.png", key, client=anon).content
            with pymupdf.open(stream=blob, filetype="png") as picture:
                self.assertEqual(picture[0].get_text().strip(), "", f"{key}: no text layer")
            self.assertIn(self.get("full-stack", "download/pdf", key, client=anon).status_code, (401, 403))

    # --- ongoing / completed follow the date --------------------------------------------

    def test_present_follows_the_date(self):
        with mock.patch("portfolio.resume_builder.date") as fake:
            fake.today.return_value = date(2030, 1, 1)         # long after the job's end date
            data = resume_builder.collect(self.fs, "https://site.example")
        self.assertNotIn("Present", json.dumps(data["experience"]))
        self.assertEqual(data["experience"][0]["period"].split(" \u2013 ")[1][-4:].isdigit(), True)
        data = resume_builder.collect(self.fs, "https://site.example")
        self.assertIn("Present", json.dumps(data["experience"]))

    def test_levels_match_the_website_bands(self):
        for rating, label in ((0, "Beginner"), (34, "Beginner"), (35, "Intermediate"), (59, "Intermediate"),
                              (60, "Advanced"), (79, "Advanced"), (80, "Expert"), (100, "Expert")):
            self.assertEqual(resume_builder.level_for(rating), label, rating)

    # --- the new fields --------------------------------------------------------------------

    def test_location_and_degree_are_validated_and_saved(self):
        ok = self.client.patch(f"{API}/tracks/{self.fs.id}", {"location": "Abuja, Nigeria"}, format="json")
        self.assertEqual(ok.status_code, 200, ok.content)
        self.assertEqual(self.client.patch(f"{API}/tracks/{self.fs.id}", {"location": "x" * 81}, format="json").status_code, 400)
        self.assertEqual(self.client.patch(f"{API}/school/{self.school.id}", {"degree": "B.Sc."}, format="json").status_code, 200)
        self.assertEqual(self.client.patch(f"{API}/school/{self.school.id}", {"degree": "x" * 101}, format="json").status_code, 400)
        self.assertEqual(self.client.patch(f"{API}/school/{self.school.id}", {"degree": ""}, format="json").status_code, 200)

    def test_duplicate_copies_location_and_degree(self):
        res = self.client.post(f"{API}/tracks/", {"username": "owner", "name": "Copy", "copy_from": self.fs.id}, format="json")
        self.assertEqual(res.status_code, 201, res.content)
        copy = PortfolioTrack.objects.get(pk=res.json()["id"])
        self.assertEqual(copy.location, "Abuja, Nigeria")
        self.assertEqual(copy.about.school.get().degree, "B.Eng. Mechanical Engineering")
        self.client.patch(f"{API}/tracks/{copy.id}", {"location": "Elsewhere"}, format="json")
        self.fs.refresh_from_db()
        self.assertEqual(self.fs.location, "Abuja, Nigeria")
