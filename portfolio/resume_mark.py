"""
The site's logo as a faint, slanted mark behind a resume: on every PDF page, and
in the header of a Word document (behind the text, like Word's own watermarks).

It is a picture only - no text - so the resume's text layer, reading order and
headings are exactly what they were. The logo is recoloured mid-grey and made
mostly see-through once, up front, so it can sit behind text without hurting it.
"""
import functools
import io
from pathlib import Path

from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt
from PIL import Image
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader

LOGO = Path(__file__).resolve().parent / "assets" / "logo-full.png"
WIDTH_MM = 150          # across the page; the picture's own padding keeps the lettering well inside it
TILT_DEGREES = 14       # slanted upwards to the right, as on the website
STRENGTH = 0.10         # how much of full opacity the grey logo keeps


@functools.lru_cache(maxsize=1)
def faint_logo():
    """(PNG bytes, width, height) of the logo, grey and mostly see-through."""
    logo = Image.open(LOGO).convert("RGBA")
    grey = Image.new("RGBA", logo.size, (110, 114, 123, 0))
    grey.putalpha(logo.getchannel("A").point(lambda alpha: int(alpha * STRENGTH)))
    out = io.BytesIO()
    grey.save(out, "PNG")
    return out.getvalue(), logo.width, logo.height


@functools.lru_cache(maxsize=1)
def _reader():
    return ImageReader(io.BytesIO(faint_logo()[0]))    # one reader, so a document embeds the picture once


def draw_on_page(canvas, page_width, page_height):
    """Draw the mark, centred and slanted, at the start of a PDF page (so everything else lies over it)."""
    _, w, h = faint_logo()
    width = WIDTH_MM * mm
    height = width * h / w
    canvas.saveState()
    canvas.translate(page_width / 2, page_height / 2)
    canvas.rotate(TILT_DEGREES)
    canvas.drawImage(_reader(), -width / 2, -height / 2, width, height, mask="auto")
    canvas.restoreState()


def add_to_docx(document):
    """Put the mark in the document's header as a picture behind the text, centred on the page."""
    section = document.sections[0]
    section.header_distance = Mm(4)                       # keeps the (empty) header line from nudging the body down
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    run.font.size = Pt(1)
    data, _, _ = faint_logo()
    run.add_picture(io.BytesIO(data), width=Mm(WIDTH_MM))

    # python-docx places pictures inline; turn this one into a floating picture behind the text
    drawing = run._r.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))
    extent = inline.find(qn("wp:extent"))
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" '
        f'behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:align>center</wp:align></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:align>center</wp:align></wp:positionV>'
        f'<wp:extent cx="{extent.get("cx")}" cy="{extent.get("cy")}"/>'
        f'<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/></wp:anchor>'
    )
    docpr = inline.find(qn("wp:docPr"))
    docpr.set("name", "Watermark")
    docpr.set("descr", "")
    anchor.append(docpr)
    for tag in ("wp:cNvGraphicFramePr", "a:graphic"):
        anchor.append(inline.find(qn(tag)))
    for xfrm in anchor.iter(qn("a:xfrm")):
        xfrm.set("rot", str(-TILT_DEGREES * 60000))      # DrawingML counts in 60,000ths of a degree; negative = anticlockwise
    drawing.remove(inline)
    drawing.append(anchor)
