"""
Builds one portfolio's resume.

collect() gathers the content once; render_pdf()/render_docx() turn it into an
ATS-friendly file (single column, real selectable text, standard fonts, plain
section headings, no tables, icons or images), and render_pages() turns the PDF
into page images for the read-only viewer, which has no text layer to copy.
"""
import hashlib
import io
import json
import re
import unicodedata
from datetime import date
from xml.sax.saxutils import escape

import pymupdf
from django.core.cache import cache
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer

# Same bands as the portfolio page (resume.jsx): the owner's 0-100 rating in words.
LEVELS = [(80, "Expert"), (60, "Advanced"), (35, "Intermediate"), (0, "Beginner")]

CACHE_SECONDS = 3600


def level_for(rating):
    return next(label for floor, label in LEVELS if rating >= floor)


def _month_year(day):
    return f"{day.month:02d}/{day.year}"


def _period(start, end, today, *, months):
    """'09/2022 - 10/2022' (or years only), with 'Present' while the end date is still ahead."""
    ongoing = end > today
    fmt = _month_year if months else (lambda d: str(d.year))
    return f"{fmt(start)} – {'Present' if ongoing else fmt(end)}"


def _bullets(text):
    """A description becomes bullet points: one per line (a leading dash/dot is ignored)."""
    lines = [re.sub(r"^[\-•*\s]+", "", line).strip() for line in (text or "").splitlines()]
    return [line for line in lines if line]


def _file_part(text):
    return re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_") or "Resume"


def collect(track, site_url):
    """Everything the resume says, from one portfolio, as plain data."""
    today = date.today()
    profile = track.profile
    user = profile.user
    about = getattr(track, "about", None)

    name = f"{user.first_name} {user.last_name}".strip() or user.username
    years = about.years_of_experience if about else 0
    headline = (about.skill if about else "") or track.name
    if years >= 1:
        headline = f"{headline} | {years}+ years of experience"

    contact = [part for part in (track.email, track.phone_number, track.location) if part]
    links = [("Portfolio", f"{site_url.rstrip('/')}/{track.slug}")]
    for label, url in (("LinkedIn", track.linkedin), ("GitHub", track.github)):
        if url:
            links.append((label, url))

    experience, education, skills, certificates = [], [], {}, []
    if about:
        for job in sorted(about.resume.all(), key=lambda r: (r.end_year, r.start_year), reverse=True):
            experience.append({
                "title": job.post, "company": job.company_name,
                "period": _period(job.start_year, job.end_year, today, months=True),
                "bullets": _bullets(job.description),
            })
            if job.certificate_link:
                certificates.append((f"{job.post}, {job.company_name}", job.certificate_link))
        for school in sorted(about.school.all(), key=lambda s: (s.end_year, s.start_year), reverse=True):
            education.append({
                "school": school.school_name, "degree": school.degree,
                "period": _period(school.start_year, school.end_year, today, months=False),
                "bullets": _bullets(school.description),
            })
            if school.certificate_link:
                certificates.append((school.school_name, school.certificate_link))
        for skill in sorted(about.proficiency.all(), key=lambda p: (-p.skill_range, p.skill_name)):
            skills.setdefault(level_for(skill.skill_range), []).append(skill.skill_name)
            if skill.certificate_link:
                certificates.append((skill.skill_name, skill.certificate_link))

    projects = [
        {"name": item.name, "category": item.category, "link": item.project_link}
        for item in track.items.all()
    ]

    return {
        "today": today.isoformat(),
        "name": name,
        "portfolio": track.name,
        "headline": headline,
        "contact": contact,
        "links": links,
        "summary": (about.description if about else "").strip(),
        "experience": experience,
        "education": education,
        # strongest first, in the order of LEVELS
        "skills": [(label, skills[label]) for _, label in LEVELS if label in skills],
        "projects": projects,
        "certificates": certificates,
    }


def file_stem(data):
    return f"{_file_part(data['name'])}_{_file_part(data['portfolio'])}_Resume"


def digest(data):
    return hashlib.sha1(json.dumps(data, sort_keys=True, default=str).encode()).hexdigest()[:16]


# ---------------------------------------------------------------- PDF --------

def _pdf_safe(text):
    """Standard PDF fonts cover Western European text. Each character outside
    that set is simplified on its own (accent dropped, e.g. a dotted O becomes
    O) or shown as ?, rather than turned into a box or an error - and the
    characters that are fine are left exactly as written."""
    text = (text or "").replace("\u2019", "'").replace("\u201c", '"').replace("\u201d", '"')
    out = []
    for char in text:
        try:
            char.encode("cp1252")
            out.append(char)
            continue
        except UnicodeEncodeError:
            pass
        base = "".join(c for c in unicodedata.normalize("NFKD", char) if not unicodedata.combining(c))
        try:
            base.encode("cp1252")
            out.append(base or "?")
        except UnicodeEncodeError:
            out.append("?")
    return "".join(out)


def _p(text):
    return escape(_pdf_safe(text))


def _link(url):
    safe = escape(_pdf_safe(url))
    return f'<link href="{safe}" color="#1a4fa0">{safe}</link>'


def render_pdf(data):
    ink, grey = colors.HexColor("#111111"), colors.HexColor("#555555")
    base = ParagraphStyle("base", fontName="Helvetica", fontSize=10, leading=13.5, textColor=ink)
    styles = {
        "name": ParagraphStyle("name", parent=base, fontName="Helvetica-Bold", fontSize=21, leading=25),
        "headline": ParagraphStyle("headline", parent=base, fontSize=11.5, leading=15, textColor=grey, spaceAfter=3),
        "small": ParagraphStyle("small", parent=base, fontSize=9.5, leading=13),
        "h": ParagraphStyle("h", parent=base, fontName="Helvetica-Bold", fontSize=10.5, leading=14, spaceBefore=9),
        "role": ParagraphStyle("role", parent=base, fontName="Helvetica-Bold", spaceBefore=4),
        "meta": ParagraphStyle("meta", parent=base, fontSize=9, textColor=grey),
        "bullet": ParagraphStyle("bullet", parent=base, leftIndent=12, bulletIndent=2, spaceBefore=1),
    }

    def heading(title):
        return [
            Paragraph(title, styles["h"]),
            HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#999999"), spaceBefore=1, spaceAfter=3),
        ]

    def bullet(text):
        return Paragraph(_p(text), styles["bullet"], bulletText="•")

    story = [Paragraph(_p(data["name"]), styles["name"]), Paragraph(_p(data["headline"]), styles["headline"])]
    if data["contact"]:
        story.append(Paragraph(_p("  |  ".join(data["contact"])), styles["small"]))
    if data["links"]:
        story.append(Paragraph("  |  ".join(f"{_p(label)}: {_link(url)}" for label, url in data["links"]), styles["small"]))

    if data["summary"]:
        story += heading("PROFESSIONAL SUMMARY") + [Paragraph(_p(data["summary"]), base)]

    if data["skills"]:
        story += heading("SKILLS")
        story += [Paragraph(f"<b>{_p(label)}:</b> {_p(', '.join(names))}", base) for label, names in data["skills"]]

    if data["experience"]:
        story += heading("PROFESSIONAL EXPERIENCE")
        for job in data["experience"]:
            block = [Paragraph(f"{_p(job['title'])} | {_p(job['company'])}", styles["role"]), Paragraph(_p(job["period"]), styles["meta"])]
            block += [bullet(b) for b in job["bullets"]]
            story.append(KeepTogether(block))

    if data["education"]:
        story += heading("EDUCATION")
        for school in data["education"]:
            title = _p(school["school"]) if not school["degree"] else f"{_p(school['degree'])} | {_p(school['school'])}"
            block = [Paragraph(title, styles["role"]), Paragraph(_p(school["period"]), styles["meta"])]
            block += [bullet(b) for b in school["bullets"]]
            story.append(KeepTogether(block))

    if data["projects"]:
        story += heading("PROJECTS")
        for project in data["projects"]:
            story.append(Paragraph(f"<b>{_p(project['name'])}</b> ({_p(project['category'])}): {_link(project['link'])}", styles["bullet"], bulletText="•"))

    if data["certificates"]:
        story += heading("CERTIFICATES")
        for label, url in data["certificates"]:
            story.append(Paragraph(f"{_p(label)}: {_link(url)}", styles["bullet"], bulletText="•"))

    story.append(Spacer(1, 2))
    buffer = io.BytesIO()
    SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=15 * mm, bottomMargin=15 * mm,
        title=f"{data['name']} - Resume", author=data["name"], subject=data["headline"],
    ).build(story)
    return buffer.getvalue()


# --------------------------------------------------------------- Word --------

def _hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run, props = OxmlElement("w:r"), OxmlElement("w:rPr")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), "1A4FA0")
    props.append(colour)
    props.append(underline)
    label = OxmlElement("w:t")
    label.text = text
    label.set(qn("xml:space"), "preserve")
    run.append(props)
    run.append(label)
    link.append(run)
    paragraph._p.append(link)


def render_docx(data):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(18)
    section.top_margin = section.bottom_margin = Mm(15)

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)
    # real Word heading styles are what ATS and screen readers key on
    heading_style = doc.styles["Heading 1"]
    heading_style.font.name, heading_style.font.size = "Calibri", Pt(11)
    heading_style.font.bold = True
    heading_style.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    heading_style.paragraph_format.space_before, heading_style.paragraph_format.space_after = Pt(10), Pt(3)

    props = doc.core_properties
    props.title, props.author, props.subject = f"{data['name']} - Resume", data["name"], data["headline"]

    title = doc.add_paragraph(style="Title")
    run = title.add_run(data["name"])
    run.font.size, run.font.bold, run.font.name = Pt(21), True, "Calibri"
    run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph(data["headline"])
    if data["contact"]:
        doc.add_paragraph("  |  ".join(data["contact"]))
    if data["links"]:
        links = doc.add_paragraph()
        for index, (label, url) in enumerate(data["links"]):
            links.add_run(("  |  " if index else "") + f"{label}: ")
            _hyperlink(links, url, url)

    def heading(text):
        doc.add_heading(text, level=1)

    def role(left, right, dates):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(5)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run(f"{left} | {right}" if right else left).bold = True
        meta = doc.add_paragraph(dates)
        meta.paragraph_format.keep_with_next = True
        meta.runs[0].font.size = Pt(9.5)
        meta.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    if data["summary"]:
        heading("PROFESSIONAL SUMMARY")
        doc.add_paragraph(data["summary"])
    if data["skills"]:
        heading("SKILLS")
        for label, names in data["skills"]:
            line = doc.add_paragraph()
            line.add_run(f"{label}: ").bold = True
            line.add_run(", ".join(names))
    if data["experience"]:
        heading("PROFESSIONAL EXPERIENCE")
        for job in data["experience"]:
            role(job["title"], job["company"], job["period"])
            for text in job["bullets"]:
                doc.add_paragraph(text, style="List Bullet")
    if data["education"]:
        heading("EDUCATION")
        for school in data["education"]:
            role(school["degree"] or school["school"], school["school"] if school["degree"] else "", school["period"])
            for text in school["bullets"]:
                doc.add_paragraph(text, style="List Bullet")
    if data["projects"]:
        heading("PROJECTS")
        for project in data["projects"]:
            line = doc.add_paragraph(style="List Bullet")
            line.add_run(project["name"]).bold = True
            line.add_run(f" ({project['category']}): ")
            _hyperlink(line, project["link"], project["link"])
    if data["certificates"]:
        heading("CERTIFICATES")
        for label, url in data["certificates"]:
            line = doc.add_paragraph(style="List Bullet")
            line.add_run(f"{label}: ")
            _hyperlink(line, url, url)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# -------------------------------------------------- read-only page images -----

def render_pages(pdf_bytes):
    """One PNG per PDF page. They carry no text layer, so there is nothing to
    select or copy; a small footer marks them as a preview copy."""
    pages = []
    with pymupdf.open(stream=pdf_bytes, filetype="pdf") as doc:
        for page in doc:
            rect = page.rect
            page.insert_text(
                pymupdf.Point(rect.width / 2 - 38, rect.height - 16), "Read-only preview",
                fontsize=7.5, color=(0.62, 0.62, 0.62),
            )
            pages.append(page.get_pixmap(dpi=130).tobytes("png"))
    return pages


# ----------------------------------------------------------- cached access ----

# looked up by name at call time (not bound here) so the renderers can be swapped out in tests
_BUILDERS = {
    "pdf": lambda data: render_pdf(data),
    "docx": lambda data: render_docx(data),
}


def cached(kind, track, data):
    """The file for this exact content (and day: 'Present' depends on it); edits change the key."""
    key = f"resume:{kind}:{track.pk}:{digest(data)}"
    if kind == "pages":
        return cache.get_or_set(key, lambda: render_pages(cached("pdf", track, data)), CACHE_SECONDS)
    return cache.get_or_set(key, lambda: _BUILDERS[kind](data), CACHE_SECONDS)
